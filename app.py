import html
import io
import json
import re
import sqlite3
import uuid
from datetime import date, datetime
from pathlib import Path

import streamlit as st

# Veilige import van AI analyzer functies met fallback
try:
    from ai.analyzer import (
        analyze_lesson,
        create_didactic_draft,
        extract_lesson_from_text,
        generate_student_goals,
        generate_spaced_retrieval,
        get_models,
        get_default_model,
        google_available
    )
except ImportError:
    from ai.analyzer import (
        analyze_lesson,
        extract_lesson_from_text,
        get_models,
    )
    def create_didactic_draft(raw_text: str, duration: int = 50, provider: str = "Google Gemini", model_name: str = "gemini-3.6-flash") -> str:
        return raw_text

    def generate_student_goals(learning_objectives: str, provider: str = "Google Gemini", model_name: str = "gemini-3.6-flash") -> list:
        lines = [line.strip().lstrip("-•123456789. ") for line in (learning_objectives or "").split("\n") if line.strip()]
        return [f"Ik kan {line}" for line in lines[:4]] or ["Ik kan de leerdoelen van deze les toepassen."]

    def generate_spaced_retrieval(lesson_text: str, provider: str = "Google Gemini", model_name: str = "gemini-3.6-flash") -> list:
        return []

    # Extra fallbacks indien niet beschikbaar in backend
    def get_default_model(provider: str = "Lokaal (Ollama)") -> str:
        if "google" in provider.lower() or "online" in provider.lower():
            return "gemini-3.6-flash"
        return "qwen2.5:7b"

    def google_available() -> bool:
        return True

try:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement, parse_xml
    from docx.oxml.ns import nsdecls, qn
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

try:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether, HRFlowable
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False

try:
    import pypdf
    PYPDF_AVAILABLE = True
except ImportError:
    PYPDF_AVAILABLE = False

# Essentiële imports voor PowerPoint-generatie via python-pptx
try:
    from pptx import Presentation
    from pptx.util import Inches as PptInches, Pt as PptPt
    from pptx.dml.color import RGBColor as PptRGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE
    PPTX_AVAILABLE = True
except ImportError:
    PPTX_AVAILABLE = False

# ============================================================
# CONFIGURATIE & DATABASE MIGRATIE
# ============================================================

APP_TITLE = "Wijze Lessen Coach"
LEERPLANNEN_DIR = Path("data/leerplannen")

old_db = Path("wijze_lessens.db")
new_db = Path("wijze_lessen.db")
if old_db.exists() and not new_db.exists():
    try:
        old_db.rename(new_db)
    except Exception:
        pass
DB_FILE = new_db

DRAFT_FILE = Path("data/draft_lesson.json")

VALID_SECTIONS = {"LESINTRO", "LESMIDDEN", "LESAFSLUITING"}

ACTIVITY_OPTIONS = [
    "Instructie", "Retrieval practice", "Individueel werk", "Duo",
    "Groepswerk", "Klassikale bespreking", "Think-Pair-Share", "Quiz",
    "Wisbordjes", "Free recall", "Lezen", "Schrijven",
    "Luisteren", "Spreken", "Anders",
]

MATERIAL_OPTIONS = [
    "Geen specifiek materiaal", "Leerboek", "Handboek / cursus",
    "Werkblad", "Tekst", "Afbeelding", "Presentatie", "Digibord",
    "Video", "Audio", "Kahoot", "Wooclap", "BookWidgets",
    "Quiz", "Wisbordjes", "Flashcards", "Anders"
]

PRINCIPLES = [
    {"number": 1, "emoji": "🧩", "short": "Voorkennis", "name": "Activeer relevante voorkennis", "description": "Activeer voorkennis die noodzakelijk is voor de nieuwe leerstof."},
    {"number": 2, "emoji": "🎯", "short": "Instructie", "name": "Geef duidelijke, gestructureerde en uitdagende instructie", "description": "Zorg voor duidelijke doelen, succescriteria, structuur en uitdaging."},
    {"number": 3, "emoji": "🔍", "short": "Voorbeelden", "name": "Gebruik voorbeelden", "description": "Gebruik concrete voorbeelden, uitgewerkte voorbeelden en modelling."},
    {"number": 4, "emoji": "🖼️", "short": "Woord & Beeld", "name": "Combineer woord en beeld", "description": "Gebruik visuele representaties wanneer die het begrip ondersteunen."},
    {"number": 5, "emoji": "⚙️", "short": "Actieve verwerking", "name": "Laat leerstof actief verwerken", "description": "Laat leerlingen informatie actief selecteren, organiseren of toepassen."},
    {"number": 6, "emoji": "🙋", "short": "Begripscontrole", "name": "Achterhaal of de hele klas het begrepen heeft", "description": "Gebruik formatieve evaluatie en gerichte checks."},
    {"number": 7, "emoji": "🪜", "short": "Scaffolding", "name": "Ondersteun bij moeilijke opdrachten", "description": "Gebruik tijdelijke scaffolding zoals stappenplannen en bouw die af."},
    {"number": 8, "emoji": "🗓️", "short": "Spacing", "name": "Spreid oefening met leerstof in de tijd", "description": "Laat leerstof op multiple momenten terugkomen (spacing)."},
    {"number": 9, "emoji": "🔀", "short": "Interleaving", "name": "Zorg voor afwisseling in oefentypes", "description": "Gebruik interleaving en verschillende probleemtypes door elkaar."},
    {"number": 10, "emoji": "📝", "short": "Toetsing als leerstrategie", "name": "Gebruik toetsing als leer- en oefenstrategie", "description": "Gebruik retrieval practice, free recall, flashcards en exit tickets."},
    {"number": 11, "emoji": "💬", "short": "Feedback", "name": "Geef feedback die leerlingen aan het denken zet", "description": "Gebruik feed-up, feedback en feed-forward."},
    {"number": 12, "emoji": "🧭", "short": "Leren leren", "name": "Leer je leerlingen effectief leren", "description": "Leer leerlingen expliciet hoe ze kunnen plannen, monitoren en evalueren."},
]

def get_principle(item_or_num):
    if isinstance(item_or_num, dict):
        raw = item_or_num.get("number", "")
    else:
        raw = item_or_num

    digits = re.findall(r'\d+', str(raw))
    target_num = int(digits[0]) if digits else None

    if target_num is not None:
        for p in PRINCIPLES:
            if p["number"] == target_num:
                return p

    return {
        "number": target_num or raw or "?",
        "emoji": "📌",
        "short": f"Bouwsteen {raw}",
        "name": f"Bouwsteen {raw}",
        "description": ""
    }

def get_principle_number(item):
    digits = re.findall(r'\d+', str(item.get("number", "")))
    return int(digits[0]) if digits else 99

def format_star_rating(score_raw):
    try:
        num = int(score_raw)
        num = max(1, min(5, num))
        stars = "⭐️" * num + "☆" * (5 - num)
        return stars, num
    except (ValueError, TypeError):
        return "☆☆☆☆☆", 0

def sanitize_pdf_text(text: str) -> str:
    if not text:
        return ""
    t = str(text)
    t = t.replace("⭐️", "*").replace("☆", "-")
    t = t.replace("—", "-").replace("–", "-")
    t = t.replace("“", '"').replace("”", '"').replace("‘", "'").replace("’", "'")
    t = t.replace("•", "*").replace("…", "...")
    cleaned = [c for c in t if ord(c) < 256]
    return html.escape("".join(cleaned))

def clean_html_breaks(text: str) -> str:
    """Verwijdert <br>, <br/>, <br /> tags (die AI soms genereert) en vervangt deze door echte newlines."""
    if not text:
        return ""
    return re.sub(r'(?i)<br\s*/?>', '\n', str(text))

def extract_text_from_file(uploaded_file) -> str:
    if not uploaded_file:
        return ""
    file_name = uploaded_file.name.lower()
    
    if file_name.endswith(".txt"):
        try:
            return uploaded_file.read().decode("utf-8", errors="ignore")
        except Exception:
            return ""
            
    elif file_name.endswith(".docx"):
        if not DOCX_AVAILABLE:
            st.error("Installeer python-docx: `pip install python-docx`")
            return ""
        try:
            doc = Document(uploaded_file)
            parts = []
            for p in doc.paragraphs:
                if p.text.strip():
                    parts.append(p.text.strip())
            for table in doc.tables:
                for row in table.rows:
                    row_vals = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_vals:
                        parts.append(" | ".join(row_vals))
            return "\n\n".join(parts)
        except Exception as e:
            st.error(f"Fout bij uitlezen Word-bestand: {e}")
            return ""
            
    elif file_name.endswith(".pdf"):
        if not PYPDF_AVAILABLE:
            st.error("Installeer pypdf: `pip install pypdf`")
            return ""
        try:
            reader = pypdf.PdfReader(uploaded_file)
            parts = []
            for idx, page in enumerate(reader.pages):
                t = page.extract_text()
                if t and t.strip():
                    parts.append(f"--- Pagina {idx + 1} ---\n" + t.strip())
            return "\n\n".join(parts)
        except Exception as e:
            st.error(f"Fout bij uitlezen PDF-bestand: {e}")
            return ""

    elif file_name.endswith(".pptx"):
        if not PPTX_AVAILABLE:
            st.error("Installeer python-pptx: `pip install python-pptx`")
            return ""
        try:
            prs = Presentation(uploaded_file)
            parts = []
            for idx, slide in enumerate(prs.slides, start=1):
                slide_texts = []
                for shape in slide.shapes:
                    if shape.has_text_frame and shape.text_frame.text.strip():
                        slide_texts.append(shape.text_frame.text.strip())
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                    nt = slide.notes_slide.notes_text_frame.text.strip()
                    if nt:
                        slide_texts.append(f"[Notities: {nt}]")
                if slide_texts:
                    parts.append(f"--- Slide {idx} ---\n" + "\n".join(slide_texts))
            return "\n\n".join(parts)
        except Exception as e:
            st.error(f"Fout bij uitlezen PowerPoint-bestand: {e}")
            return ""
            
    return ""

def load_curriculum_goals(net="GO!", graad="3", finaliteit="Dubbele finaliteit", vak="Nederlands"):
    if not LEERPLANNEN_DIR.exists():
        return []
    for file in LEERPLANNEN_DIR.rglob("*.json"):
        try:
            with open(file, "r", encoding="utf-8") as f:
                data = json.load(f)
                meta = data.get("metadata", {})
                if (meta.get("net") == net and str(meta.get("graad")) == str(graad) and meta.get("finaliteit") == finaliteit and meta.get("vak") == vak):
                    return data.get("doelen", [])
        except Exception:
            continue
    return []

def create_phase_id(): 
    return uuid.uuid4().hex

def empty_phase(section="LESMIDDEN", phase_number=1):
    return {
        "phase_id": create_phase_id(), "section": section, 
        "phase_number": phase_number if section == "LESMIDDEN" else None, 
        "title": "" if section == "LESMIDDEN" else section, 
        "lesson_objectives": "", "learning_content": "", "organization": "", 
        "duration": 5, "teacher_activity": "", "student_activity": "", 
        "activity_type": [], "materials": []
    }

def normalize_multi_value(value):
    if value is None: return []
    if isinstance(value, (list, tuple)): return [str(item) for item in value if str(item).strip()]
    value = str(value).strip()
    if not value: return []
    try:
        parsed = json.loads(value)
        if isinstance(parsed, list): return [str(item) for item in parsed if str(item).strip()]
    except (json.JSONDecodeError, TypeError): pass
    return [value]

def serialize_multi_value(values): 
    return json.dumps(values or [], ensure_ascii=False)

def format_multi_value(value):
    values = normalize_multi_value(value)
    return ", ".join(values) if values else "Geen"

def normalize_phase(phase, section=None, phase_number=None):
    if hasattr(phase, "keys"): phase = dict(phase)
    elif not isinstance(phase, dict): phase = {}

    phase_id = phase.get("phase_id") or phase.get("_phase_id") or create_phase_id()
    inferred_section = phase.get("section") or section or "LESMIDDEN"
    if inferred_section not in VALID_SECTIONS: inferred_section = "LESMIDDEN"

    raw_number = phase.get("phase_number", phase_number)
    try: normalized_number = int(raw_number) if raw_number is not None else None
    except (TypeError, ValueError): normalized_number = phase_number
    if inferred_section != "LESMIDDEN": normalized_number = None

    try: duration = int(phase.get("duration", phase.get("duur", phase.get("tijd", 5))) or 5)
    except (TypeError, ValueError): duration = 5

    # Filter hier direct eventuele <br> tags weg
    lesson_obj = clean_html_breaks(phase.get("lesson_objectives") or phase.get("lesdoel") or phase.get("lesdoelen") or phase.get("doelen") or "")
    content = clean_html_breaks(phase.get("learning_content") or phase.get("leerinhoud") or phase.get("leerinhouden") or phase.get("inhoud") or "")
    org = clean_html_breaks(phase.get("lesverloop") or phase.get("organization") or phase.get("organisatie") or phase.get("verloop") or phase.get("lesson_flow") or "")
    teacher = clean_html_breaks(phase.get("teacher_activity") or phase.get("leraar") or phase.get("leerkracht") or phase.get("docent") or "")
    student = clean_html_breaks(phase.get("student_activity") or phase.get("leerlingen") or phase.get("leerling") or phase.get("student") or "")
    
    act_types_raw = normalize_multi_value(phase.get("activity_type") or phase.get("werkvorm") or phase.get("werkvormen"))
    activity_types = []
    for raw in act_types_raw:
        match = next((opt for opt in ACTIVITY_OPTIONS if opt.lower() in raw.lower()), None)
        if match and match not in activity_types:
            activity_types.append(match)

    mat_raw = normalize_multi_value(phase.get("materials") or phase.get("media") or phase.get("materialen") or phase.get("leermiddelen"))
    materials = []
    for raw in mat_raw:
        match = next((opt for opt in MATERIAL_OPTIONS if opt.lower() in raw.lower()), None)
        if match and match not in materials:
            materials.append(match)

    default_title = inferred_section if inferred_section != "LESMIDDEN" else f"Lesfase {normalized_number or 1}"
    title_str = str(phase.get("title") or phase.get("titel") or phase.get("naam") or default_title)

    return {
        "phase_id": phase_id, "section": inferred_section, "phase_number": normalized_number, "title": title_str,
        "lesson_objectives": str(lesson_obj), "learning_content": str(content), "organization": str(org),
        "duration": max(1, min(duration, 120)), "teacher_activity": str(teacher), "student_activity": str(student),
        "activity_type": activity_types, "materials": materials
    }

def normalize_phase_list(phases, default_middle_count=3):
    if not isinstance(phases, list): phases = []
    intro, closing, middle = None, None, []
    for raw in phases:
        p = normalize_phase(raw)
        if p["section"] == "LESINTRO" and intro is None: intro = p
        elif p["section"] == "LESAFSLUITING" and closing is None: closing = p
        else: middle.append(p)

    if intro is None: intro = empty_phase("LESINTRO")
    if closing is None: closing = empty_phase("LESAFSLUITING")
    if not middle:
        count = max(1, int(default_middle_count or 1))
        middle = [empty_phase("LESMIDDEN", n) for n in range(1, count + 1)]

    for n, p in enumerate(middle, start=1): p["section"], p["phase_number"], p["title"] = "LESMIDDEN", n, f"Lesfase {n}"
    intro["section"], intro["phase_number"], intro["title"] = "LESINTRO", None, "LESINTRO"
    closing["section"], closing["phase_number"], closing["title"] = "LESAFSLUITING", None, "LESAFSLUITING"
    return [intro, *middle, closing]

def renumber_middle_phases(phases):
    if not isinstance(phases, list): phases = []
    normalized = [normalize_phase(p) for p in phases]
    intro, closing, middle = None, None, []
    for p in normalized:
        if p["section"] == "LESINTRO":
            if intro is None: intro = p
        elif p["section"] == "LESAFSLUITING":
            if closing is None: closing = p
        else:
            p["section"] = "LESMIDDEN"
            middle.append(p)

    if intro is None: intro = empty_phase("LESINTRO")
    if closing is None: closing = empty_phase("LESAFSLUITING")
    if not middle: middle.append(empty_phase("LESMIDDEN", 1))

    for n, p in enumerate(middle, start=1): p["section"], p["phase_number"], p["title"] = "LESMIDDEN", n, f"Lesfase {n}"
    intro["section"], intro["phase_number"], intro["title"] = "LESINTRO", None, "LESINTRO"
    closing["section"], closing["phase_number"], closing["title"] = "LESAFSLUITING", None, "LESAFSLUITING"
    return [intro, *middle, closing]

def calculate_phase_duration(phases):
    normalized = normalize_phase_list(phases, default_middle_count=1)
    return sum(int(p["duration"]) for p in normalized)

def show_duration_check(total_duration, phases):
    phase_total = calculate_phase_duration(phases)
    try: total_duration = int(total_duration)
    except (TypeError, ValueError): total_duration = 0

    difference = total_duration - phase_total
    col1, col2, col3 = st.columns(3)
    with col1: st.metric("Ingestelde lesduur", f"{total_duration} min.")
    with col2: st.metric("Som van de fases", f"{phase_total} min.")
    with col3: st.metric("Verschil", f"{difference:+d} min.")

    if difference == 0: st.success("De totale lesduur komt exact overeen met de som van de lesfasen.")
    elif difference > 0: st.warning(f"De ingestelde lesduur is {difference} minuut langer dan de som van de fases.")
    else: st.warning(f"De fases duren samen {abs(difference)} minuut langer dan de ingestelde totale lesduur.")

def phase_has_content(phase):
    phase = normalize_phase(phase)
    text_fields = [phase["lesson_objectives"], phase["learning_content"], phase["organization"], phase["teacher_activity"], phase["student_activity"]]
    if any(str(v).strip() for v in text_fields): return True
    if phase["activity_type"] or phase["materials"]: return True
    return False

def merge_text_value(current, extracted):
    current, extracted = str(current) if current is not None else "", str(extracted) if extracted is not None else ""
    return extracted if extracted.strip() else current

def merge_phase(current, extracted):
    current = normalize_phase(current)
    extracted = normalize_phase(extracted, section=current["section"], phase_number=current["phase_number"])
    merged = dict(current)
    merged["lesson_objectives"] = merge_text_value(current["lesson_objectives"], extracted["lesson_objectives"])
    merged["learning_content"] = merge_text_value(current["learning_content"], extracted["learning_content"])
    merged["organization"] = merge_text_value(current["organization"], extracted["organization"])
    merged["teacher_activity"] = merge_text_value(current["teacher_activity"], extracted["teacher_activity"])
    merged["student_activity"] = merge_text_value(current["student_activity"], extracted["student_activity"])
    if extracted["activity_type"]: merged["activity_type"] = extracted["activity_type"]
    if extracted["materials"]: merged["materials"] = extracted["materials"]
    raw_duration = extracted.get("duration") if isinstance(extracted, dict) else None
    if raw_duration is not None:
        try:
            extracted_duration = int(raw_duration)
            if 1 <= extracted_duration <= 120: merged["duration"] = extracted_duration
        except (TypeError, ValueError): pass
    return merged

def merge_ai_phases_into_current(current_phases, extracted_phases):
    current = normalize_phase_list(current_phases, default_middle_count=1)
    if not extracted_phases or not isinstance(extracted_phases, list): return current

    valid_extracted = [p for p in extracted_phases if phase_has_content(p)]
    if not valid_extracted: return current

    current_intro = next(p for p in current if p["section"] == "LESINTRO")
    current_closing = next(p for p in current if p["section"] == "LESAFSLUITING")

    if len(valid_extracted) == 1:
        current_middle = [merge_phase(empty_phase("LESMIDDEN", 1), valid_extracted[0])]
    elif len(valid_extracted) == 2:
        current_intro = merge_phase(current_intro, valid_extracted[0])
        current_middle = [merge_phase(empty_phase("LESMIDDEN", 1), valid_extracted[1])]
    else:
        current_intro = merge_phase(current_intro, valid_extracted[0])
        current_closing = merge_phase(current_closing, valid_extracted[-1])
        middle_extracted = valid_extracted[1:-1]
        current_middle = [merge_phase(empty_phase("LESMIDDEN", idx), phase) for idx, phase in enumerate(middle_extracted, start=1)]

    return renumber_middle_phases([current_intro, *current_middle, current_closing])

def apply_extracted_lesson_to_session(extracted, available_goals=None):
    if not isinstance(extracted, dict): raise ValueError("De AI gaf geen geldig lesobject terug.")
    
    for wrapper in ["lesplan", "lesvoorbereiding", "les", "lesson", "data", "schema", "output"]:
        if wrapper in extracted and isinstance(extracted[wrapper], dict):
            extracted = extracted[wrapper]
            break
            
    if len(extracted) == 1:
        first_val = list(extracted.values())[0]
        if isinstance(first_val, dict):
            extracted = first_val

    title_val = extracted.get("title") or extracted.get("titel") or extracted.get("lesnaam") or ""
    class_val = extracted.get("class_name") or extracted.get("klas") or ""
    topic_val = extracted.get("topic") or extracted.get("onderwerp") or extracted.get("thema") or ""
    
    # Filter hier eventuele <br> tags weg
    objectives_val = clean_html_breaks(extracted.get("learning_objectives") or extracted.get("algemene_leerdoelen") or extracted.get("lesdoelen") or extracted.get("doelen") or "")
    notes_val = clean_html_breaks(extracted.get("notes") or extracted.get("opmerkingen") or "")

    if title_val: st.session_state["new_title"] = str(title_val)
    if class_val: st.session_state["new_class_name"] = str(class_val)
    if topic_val: st.session_state["new_topic"] = str(topic_val)
    if objectives_val: st.session_state["new_learning_objectives"] = str(objectives_val)
    if notes_val: st.session_state["new_notes"] = str(notes_val)

    raw_dur = extracted.get("duration") or extracted.get("duur") or extracted.get("lesduur") or extracted.get("tijd")
    if raw_dur is not None:
        try:
            num_match = re.search(r'\d+', str(raw_dur))
            if num_match: st.session_state["new_duration"] = int(num_match.group())
        except Exception: pass

    ai_goals = extracted.get("curriculum_goals") or extracted.get("leerplandoelen") or []
    if available_goals and ai_goals:
        matched_options = [f"{g['code']} - {g['omschrijving']} ({g.get('cluster', '')})" for g in available_goals if g["code"] in ai_goals or any(g["code"] in str(item) for item in ai_goals)]
        st.session_state["new_curriculum_goals"] = matched_options
        st.session_state["selected_goals_widget"] = matched_options

    current_phases = st.session_state.get("new_phases", normalize_phase_list([], default_middle_count=3))
    
    raw_phases = (
        extracted.get("phases")
        or extracted.get("lesfasen")
        or extracted.get("fases")
        or extracted.get("fasen")
        or extracted.get("lesfases")
        or extracted.get("stappen")
        or extracted.get("lesverloop")
        or extracted.get("activiteiten")
        or []
    )
    if isinstance(raw_phases, dict):
        raw_phases = list(raw_phases.values())

    if not raw_phases:
        for k, v in extracted.items():
            if isinstance(v, list) and len(v) > 0 and isinstance(v[0], dict):
                if any(f in v[0] for f in ["title", "titel", "section", "inhoud", "organization", "verloop", "duur", "duration", "lesson_objectives", "learning_content"]):
                    raw_phases = v
                    break

    new_phases = merge_ai_phases_into_current(current_phases, raw_phases)
    st.session_state["new_phases"] = new_phases

    if not raw_phases:
        st.session_state["new_ai_warning"] = "⚠️ De AI gaf geen lesfases terug. De algemene gegevens zijn ingevuld — vul de fases hieronder aan of probeer opnieuw."
    elif not any(phase_has_content(p) for p in raw_phases):
        st.session_state["new_ai_warning"] = "⚠️ De AI-fases bevatten geen bruikbare inhoud. Vul de lesfases hieronder handmatig aan of probeer opnieuw."
    else:
        st.session_state.pop("new_ai_warning", None)

    for phase in new_phases:
        base = f"new_new_{phase['phase_id']}"
        st.session_state[f"{base}_obj"] = phase["lesson_objectives"]
        st.session_state[f"{base}_content"] = phase["learning_content"]
        st.session_state[f"{base}_org"] = phase["organization"]
        st.session_state[f"{base}_activity"] = phase["activity_type"]
        st.session_state[f"{base}_materials"] = phase["materials"]
        st.session_state[f"{base}_time"] = phase["duration"]

def initialize_new_lesson_state():
    defaults = {
        "new_title": "", "new_class_name": "", "new_grade": "3", "new_study_direction": "Dubbele finaliteit",
        "new_topic": "", "new_learning_objectives": "", "new_notes": "", "new_duration": 50,
        "new_phases": normalize_phase_list([], default_middle_count=3), "new_ai_text": "", "new_ai_extraction_done": False,
        "new_network": "GO!", "new_subject": "Nederlands", "new_curriculum_goals": [], "new_extra_prompt": "",
        "_last_uploaded_file": "", "new_ai_warning": None
    }
    
    # 1. Zorg dat ALLE keys bestaan om KeyError te voorkomen
    for key, value in defaults.items():
        if key not in st.session_state:
            st.session_state[key] = value

    # 2. Als we al een keer geïnitialiseerd zijn, stop hier
    if st.session_state.get("_state_initialized"):
        return

    # 3. Probeer een concept te laden
    if DRAFT_FILE.exists():
        try:
            with open(DRAFT_FILE, "r", encoding="utf-8") as f:
                draft = json.load(f)
                if draft and isinstance(draft, dict):
                    st.session_state["new_title"] = draft.get("title", "")
                    st.session_state["new_class_name"] = draft.get("class_name", "")
                    st.session_state["new_grade"] = draft.get("grade", "3")
                    st.session_state["new_study_direction"] = draft.get("study_direction", "Dubbele finaliteit")
                    st.session_state["new_topic"] = draft.get("topic", "")
                    st.session_state["new_learning_objectives"] = clean_html_breaks(draft.get("learning_objectives", ""))
                    st.session_state["new_notes"] = clean_html_breaks(draft.get("notes", ""))
                    st.session_state["new_duration"] = draft.get("duration", 50)
                    st.session_state["new_network"] = draft.get("network", "GO!")
                    st.session_state["new_subject"] = draft.get("subject", "Nederlands")
                    st.session_state["new_curriculum_goals"] = draft.get("curriculum_goals", [])
                    st.session_state["selected_goals_widget"] = draft.get("curriculum_goals", [])
                    st.session_state["new_phases"] = draft.get("phases", normalize_phase_list([], default_middle_count=3))
                    date_str = draft.get("date")
                    if date_str:
                        try: st.session_state["new_date"] = datetime.strptime(date_str, "%Y-%m-%d").date()
                        except ValueError: pass
        except Exception: pass
    
    st.session_state["_state_initialized"] = True

def reset_new_lesson():
    keys = ["new_title", "new_class_name", "new_grade", "new_study_direction", "new_topic", "new_learning_objectives", "new_notes", "new_duration", "new_phases", "new_ai_text", "new_ai_extraction_done", "new_date", "new_curriculum_goals", "selected_goals_widget", "new_extra_prompt", "_last_uploaded_file", "new_ai_warning", "_state_initialized"]
    for key in keys: st.session_state.pop(key, None)
    if DRAFT_FILE.exists():
        try: DRAFT_FILE.unlink()
        except Exception: pass

def format_for_smartschool(lesson, phases) -> str:
    lines = [f"📚 LESFICHE: {lesson['title']}", f"Vak: {lesson['subject'] or 'Nederlands'} | Klas: {lesson['class_name'] or 'Klas'} | Duur: {lesson['duration'] or 50} min.", "-" * 45, "\n🎯 LESDOELEN (Wat leren we vandaag?):", lesson['learning_objectives'] or "Zie lesverloop."]
    if lesson['curriculum_goals']:
        lines.append("\n📌 GEKOPPELDE LEERPLANDOELEN:")
        for g in normalize_multi_value(lesson['curriculum_goals']): lines.append(f"• {g}")
    lines.append("\n⏱️ LESVERLOOP & ACTIVITEITEN:")
    normalized = normalize_phase_list([dict(p) if hasattr(p, "keys") else p for p in phases], default_middle_count=1)
    for idx, raw_phase in enumerate(normalized, start=1):
        phase = normalize_phase(raw_phase)
        label = phase["section"] if phase["section"] in ["LESINTRO", "LESAFSLUITING"] else f"Fase {phase['phase_number'] or idx}"
        lines.append(f"\n[{label}] ({phase['duration']} min.)")
        if phase['learning_content']: lines.append(f"• Inhoud: {phase['learning_content']}")
        if phase['organization']: lines.append(f"• Verloop: {phase['organization']}")
        if phase['materials']: lines.append(f"• Media: {format_multi_value(phase['materials'])}")
    if lesson['notes']: lines.append(f"\n📝 OPMERKINGEN / MATERIAAL:\n{lesson['notes']}")
    return "\n".join(lines)

def export_analysis_to_pdf(lesson, analysis, version_str="Actuele versie") -> io.BytesIO:
    if not PDF_AVAILABLE or not analysis: return None
    bio = io.BytesIO()
    doc = SimpleDocTemplate(bio, pagesize=A4, leftMargin=1.5 * cm, rightMargin=1.5 * cm, topMargin=1.5 * cm, bottomMargin=1.5 * cm)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('DocTitle', parent=styles['Heading1'], fontName='Helvetica-Bold', fontSize=18, leading=22, textColor=colors.HexColor("#1E3A8A"), spaceAfter=4)
    sub_style = ParagraphStyle('DocSub', parent=styles['Normal'], fontName='Helvetica-Oblique', fontSize=9.5, leading=13, textColor=colors.HexColor("#4B5563"), spaceAfter=12)
    h2_style = ParagraphStyle('Heading2_Custom', parent=styles['Heading2'], fontName='Helvetica-Bold', fontSize=12, leading=16, textColor=colors.HexColor("#1E3A8A"), spaceBefore=10, spaceAfter=6)
    body_style = ParagraphStyle('Body_Custom', parent=styles['Normal'], fontName='Helvetica', fontSize=9, leading=13, textColor=colors.HexColor("#1F2937"), spaceAfter=4)
    bold_body = ParagraphStyle('BoldBody', parent=body_style, fontName='Helvetica-Bold')
    story = [Paragraph(f"Didactische Analyse: {sanitize_pdf_text(lesson.get('title', 'Les'))}", title_style), Paragraph("Wijze Lessen Coach &bull; Pedagogische Evaluatie", sub_style)]
    analyzed_time = analysis.get("analyzed_at", "Niet geregistreerd")
    meta_data = [
        [Paragraph("<b>Vak & Klas:</b>", body_style), Paragraph(sanitize_pdf_text(f"{lesson.get('subject', 'Nederlands')} | {lesson.get('class_name', '—')} ({lesson.get('grade', '3')}e graad)"), body_style)],
        [Paragraph("<b>Lesversie:</b>", body_style), Paragraph(sanitize_pdf_text(version_str), bold_body)],
        [Paragraph("<b>Tijdstip AI-analyse:</b>", body_style), Paragraph(sanitize_pdf_text(analyzed_time), bold_body)],
        [Paragraph("<b>Totale Score:</b>", body_style), Paragraph(sanitize_pdf_text(f"{analysis.get('total_score', '—')} / 60"), bold_body)],
    ]
    meta_table = Table(meta_data, colWidths=[4.2 * cm, 13.8 * cm])
    meta_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor("#F3F4F6")),
        ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor("#D1D5DB")),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#E5E7EB")),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('LEFTPADDING', (0, 0), (-1, -1), 8),
        ('RIGHTPADDING', (0, 0), (-1, -1), 8),
    ]))
    story.append(meta_table)
    story.append(Spacer(1, 10))
    if analysis.get("summary"):
        story.append(Paragraph("Samenvatting Didactische Kwaliteit", h2_style))
        story.append(Paragraph(sanitize_pdf_text(analysis["summary"]), body_style))
        story.append(Spacer(1, 6))
    if analysis.get("priorities"):
        story.append(Paragraph("Belangrijkste Prioriteiten", h2_style))
        for idx, prio in enumerate(analysis["priorities"], start=1): story.append(Paragraph(f"<b>{idx}.</b> {sanitize_pdf_text(prio)}", body_style))
        story.append(Spacer(1, 8))
    principles = analysis.get("principles", [])
    if principles:
        story.append(Paragraph("Overzicht Didactische Rapportkaart", h2_style))
        principles_sorted = sorted(principles, key=get_principle_number)
        overview_data = [[Paragraph("<b>Bouwsteen</b>", bold_body), Paragraph("<b>Status</b>", bold_body), Paragraph("<b>Score</b>", bold_body)]]
        for item in principles_sorted:
            p_obj = get_principle(item)
            stars, score_num = format_star_rating(item.get("score", 0))
            stars_ascii = "*" * score_num + "-" * (5 - score_num)
            overview_data.append([Paragraph(sanitize_pdf_text(f"B{p_obj['number']}. {p_obj['name']}"), body_style), Paragraph(sanitize_pdf_text(str(item.get('status', 'Beoordeeld'))), body_style), Paragraph(sanitize_pdf_text(f"{score_num}/5 ({stars_ascii})"), body_style)])
        overview_table = Table(overview_data, colWidths=[10.5 * cm, 4.0 * cm, 3.5 * cm])
        overview_table.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#1E3A8A")), ('TEXTCOLOR', (0, 0), (-1, 0), colors.white), ('BOTTOMPADDING', (0, 0), (-1, -1), 3.5), ('TOPPADDING', (0, 0), (-1, -1), 3.5), ('LEFTPADDING', (0, 0), (-1, -1), 6), ('RIGHTPADDING', (0, 0), (-1, -1), 6), ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#D1D5DB"))]))
        story.append(overview_table)
        story.append(Spacer(1, 14))
        story.append(Paragraph("Gedetailleerde Evaluatie & Adviezen per Bouwsteen", h2_style))
        for item in principles_sorted:
            p_obj = get_principle(item)
            stars, score_num = format_star_rating(item.get("score", 0))
            block_elements = [Paragraph(sanitize_pdf_text(f"<b>B{p_obj['number']}. {p_obj['name']}</b> &mdash; Score: {score_num}/5 [{item.get('status', 'Beoordeeld')}]"), bold_body), Paragraph(f"<b>Bewijs uit de les:</b> {sanitize_pdf_text(item.get('evidence', 'Geen specifiek bewijs.'))}", body_style)]
            sug1 = item.get("suggestion_1") or item.get("suggestie_1")
            sug2 = item.get("suggestion_2") or item.get("suggestie_2")
            sug3 = item.get("suggestion_3") or item.get("suggestie_3")
            if sug1 or sug2 or sug3:
                block_elements.append(Paragraph("<b>Verbeteradviezen:</b>", body_style))
                if sug1: block_elements.append(Paragraph(f"&bull; 1. {sanitize_pdf_text(sug1)}", body_style))
                if sug2: block_elements.append(Paragraph(f"&bull; 2. {sanitize_pdf_text(sug2)}", body_style))
                if sug3: block_elements.append(Paragraph(f"&bull; 3. {sanitize_pdf_text(sug3)}", body_style))
            block_elements.append(Spacer(1, 6))
            story.append(KeepTogether(block_elements))
    doc.build(story)
    bio.seek(0)
    return bio

def export_lesson_to_docx(lesson, phases, analysis=None) -> io.BytesIO:
    if not DOCX_AVAILABLE: return None
    doc = Document()
    for section in doc.sections: section.top_margin, section.bottom_margin, section.left_margin, section.right_margin = Inches(0.8), Inches(0.8), Inches(0.8), Inches(0.8)
    title_p = doc.add_paragraph()
    title_run = title_p.add_run(f"Lesvoorbereiding: {lesson['title']}")
    title_run.font.size, title_run.font.bold, title_run.font.color.rgb = Pt(18), True, RGBColor(30, 58, 138)
    title_p.paragraph_format.space_after = Pt(4)
    sub_p = doc.add_paragraph()
    sub_run = sub_p.add_run(f"Vak: {lesson['subject'] or 'Nederlands'}  |  Klas: {lesson['class_name'] or '—'}  |  Datum: {lesson['date'] or '—'}  |  Duur: {lesson['duration'] or 50} min.")
    sub_run.font.size, sub_run.font.italic = Pt(10), True
    sub_p.paragraph_format.space_after = Pt(14)
    doc.add_heading("1. Context & Doelstellingen", level=2)
    p_meta = doc.add_paragraph()
    p_meta.add_run("• Onderwijsnet: ").bold = True; p_meta.add_run(f"{lesson['network'] or 'GO!'}\n")
    p_meta.add_run("• Graad & Finaliteit: ").bold = True; p_meta.add_run(f"{lesson['grade'] or '3'}e graad - {lesson['study_direction'] or 'Dubbele finaliteit'}\n")
    p_meta.add_run("• Onderwerp: ").bold = True; p_meta.add_run(f"{lesson['topic'] or '—'}")
    p_meta.paragraph_format.space_after = Pt(8)
    p_doelen = doc.add_paragraph()
    p_doelen.add_run("Algemene lesdoelen:\n").bold = True; p_doelen.add_run(lesson['learning_objectives'] or "Geen algemene lesdoelen opgegeven.")
    p_doelen.paragraph_format.space_after = Pt(8)
    if lesson['curriculum_goals']:
        p_lpd = doc.add_paragraph()
        p_lpd.add_run("Gekoppelde officiële leerplandoelen:\n").bold = True
        for g in normalize_multi_value(lesson['curriculum_goals']): p_lpd.add_run(f"  ✔ {g}\n")
        p_lpd.paragraph_format.space_after = Pt(12)
    doc.add_heading("2. Lesstructuur & Verloop", level=2)
    normalized = normalize_phase_list([dict(p) if hasattr(p, "keys") else p for p in phases], default_middle_count=1)
    table = doc.add_table(rows=1, cols=5)
    table.alignment, table.autofit = WD_TABLE_ALIGNMENT.CENTER, False
    col_widths = [Inches(1.1), Inches(1.5), Inches(1.5), Inches(2.2), Inches(0.7)]
    hdr_titles = ["Fase", "Lesdoelen", "Leerinhoud", "Lesverloop, Werkvorm & Media", "Tijd"]
    for i, title in enumerate(hdr_titles):
        table.rows[0].cells[i].text = title; table.rows[0].cells[i].width = col_widths[i]
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="1E3A8A"/>')
        table.rows[0].cells[i]._tc.get_or_add_tcPr().append(shading)
        for p in table.rows[0].cells[i].paragraphs:
            for r in p.runs: r.font.bold, r.font.size, r.font.color.rgb = True, Pt(9.5), RGBColor(255, 255, 255)
    for raw_phase in normalized:
        phase = normalize_phase(raw_phase)
        row_cells = table.add_row().cells
        label = phase["section"] if phase["section"] in ["LESINTRO", "LESAFSLUITING"] else f"LESMIDDEN\nFase {phase['phase_number'] or ''}"
        verloop_text = phase["organization"] or "—"
        extra_info = []
        if phase["activity_type"]: extra_info.append(f"Werkvorm: {format_multi_value(phase['activity_type'])}")
        if phase["materials"]: extra_info.append(f"Media: {format_multi_value(phase['materials'])}")
        if phase["teacher_activity"]: extra_info.append(f"Leraar: {phase['teacher_activity']}")
        if phase["student_activity"]: extra_info.append(f"Leerlingen: {phase['student_activity']}")
        if extra_info: verloop_text += "\n\n" + "\n".join(extra_info)
        row_data = [label, phase["lesson_objectives"] or "—", phase["learning_content"] or "—", verloop_text, f"{phase['duration']} min."]
        for i, text in enumerate(row_data):
            row_cells[i].text = text; row_cells[i].width = col_widths[i]
            for p in row_cells[i].paragraphs:
                for r in p.runs: r.font.size = Pt(8.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(14)
    if analysis:
        doc.add_heading("3. Didactische Analyse (Wijze Lessen Coach)", level=2)
        if analysis.get("total_score") is not None:
            p_s = doc.add_paragraph()
            p_s.add_run(f"Totale score: {analysis['total_score']} / 60\n").bold = True
            if analysis.get("summary"): p_s.add_run(f"Samenvatting: {analysis['summary']}\n")
            if analysis.get("priorities"):
                p_prio = doc.add_paragraph()
                p_prio.add_run("Belangrijkste prioriteiten:\n").bold = True
                for i, p in enumerate(analysis["priorities"], start=1): p_prio.add_run(f"  {i}. {p}\n")
        if analysis.get("principles"):
            doc.add_heading("Evaluatie per bouwsteen", level=3)
            for item in sorted(analysis["principles"], key=get_principle_number):
                p_obj = get_principle(item)
                stars, score_num = format_star_rating(item.get("score", 0))
                p_b = doc.add_paragraph()
                p_b.add_run(f"{p_obj['emoji']} B{p_obj['number']}. {p_obj['name']} — {stars} ({score_num}/5)\n").bold = True
                p_b.add_run(f"• Bewijs uit de les: {item.get('evidence', '—')}\n")
                sug1 = item.get("suggestion_1") or item.get("suggestie_1") or item.get("advies_1")
                sug2 = item.get("suggestion_2") or item.get("suggestie_2") or item.get("advies_2")
                sug3 = item.get("suggestion_3") or item.get("suggestie_3") or item.get("advies_3")
                p_b.add_run("• Verbeteradviezen:\n").italic = True
                if sug1: p_b.add_run(f"   1. {sug1}\n")
                if sug2: p_b.add_run(f"   2. {sug2}\n")
                if sug3: p_b.add_run(f"   3. {sug3}\n")
                p_b.paragraph_format.space_after = Pt(6)
    bio = io.BytesIO(); doc.save(bio); bio.seek(0); return bio

def export_lesson_to_text(lesson, phases, analysis=None) -> str:
    lines = ["=" * 70, f"LESVOORBEREIDING: {lesson['title']}", "=" * 70, "", "1. BASISINFORMATIE", "-" * 30]
    lines.append(f"• Datum:                      {lesson['date'] or 'Niet opgegeven'}")
    lines.append(f"• Klas:                       {lesson['class_name'] or 'Niet opgegeven'}")
    lines.append(f"• Graad / Leerjaar:           {lesson['grade'] or 'Niet opgegeven'}")
    lines.append(f"• Studierichting / Finaliteit: {lesson['study_direction'] or 'Niet opgegeven'}")
    lines.append(f"• Onderwijsnet:               {lesson['network'] or 'Niet opgegeven'}")
    lines.append(f"• Vak:                        {lesson['subject'] or 'Niet opgegeven'}")
    lines.append(f"• Onderwerp:                  {lesson['topic'] or 'Niet opgegeven'}")
    lines.append(f"• Totale lesduur:             {lesson['duration'] or 0} minuten\n")
    lines.extend(["2. DOELSTELLINGEN", "-" * 30, "ALGEMENE LESDOELEN:", lesson['learning_objectives'] or "Geen algemene lesdoelen.", ""])
    if lesson['curriculum_goals']:
        lines.append("GEKOPPELDE OFFICIELE LEERPLANDOELEN:")
        for g_str in normalize_multi_value(lesson["curriculum_goals"]): lines.append(f"  [x] {g_str}")
        lines.append("")
    if lesson['notes']: lines.extend(["ALGEMENE OPMERKINGEN:", lesson['notes'], ""])
    lines.extend(["=" * 70, "3. LESSTRUCTUUR & VERLOOP", "=" * 70])
    for index, raw_phase in enumerate(normalize_phase_list([dict(p) if hasattr(p, "keys") else p for p in phases], default_middle_count=1), start=1):
        phase = normalize_phase(raw_phase)
        label = phase["section"] if phase["section"] in ["LESINTRO", "LESAFSLUITING"] else f"LESMIDDEN - Lesfase {phase['phase_number'] or index}"
        lines.extend(["", f">>> {label} ({phase['duration']} min.) <<<", f"• Lesdoelen:               {phase['lesson_objectives'] or '—'}", f"• Leerinhouden:            {phase['learning_content'] or '—'}", f"• Werkvormen:              {format_multi_value(phase['activity_type'])}", f"• Media / Leermiddelen:    {format_multi_value(phase['materials'])}", "• Lesverloop & Organisatie:", f"  {phase['organization'] or '—'}"])
        if phase['teacher_activity']: lines.append(f"• Leraar activiteit:       {phase['teacher_activity']}")
        if phase['student_activity']: lines.append(f"• Leerling activiteit:     {phase['student_activity']}")
        lines.append("-" * 50)
    if analysis:
        lines.extend(["", "=" * 70, "4. AI-DIDACTISCHE ANALYSE (12 BOUWSTENEN)", "=" * 70])
        if analysis.get("total_score") is not None: lines.append(f"TOTALE SCORE: {analysis['total_score']} / 60")
        if analysis.get("summary"): lines.append(f"SAMENVATTING: {analysis['summary']}")
        lines.extend(["", "RAPPORTKAART BOUWSTENEN:", "-" * 30])
        for item in sorted(analysis.get("principles", []), key=get_principle_number):
            p_obj = get_principle(item)
            stars, score_num = format_star_rating(item.get("score", 0))
            lines.append(f"{p_obj['emoji']} B{p_obj['number']}. {p_obj['short']:<20} {stars} ({score_num}/5)")
        if analysis.get("priorities"):
            lines.extend(["", "BELANGRIJKSTE PRIORITEITEN / WERKPUNTEN:"])
            for i, p in enumerate(analysis["priorities"], start=1): lines.append(f"  {i}. {p}")
        if analysis.get("principles"):
            lines.extend(["", "GEDETAILLEERDE EVALUATIE PER BOUWSTEEN:", "-" * 40])
            for item in sorted(analysis["principles"], key=get_principle_number):
                p_obj = get_principle(item)
                stars, score_num = format_star_rating(item.get("score", 0))
                lines.extend([f"\n[{p_obj['emoji']} B{p_obj['number']}. {p_obj['name']}] — Score: {stars} ({score_num}/5) [{item.get('status', '')}]", f"• Bewijs uit de les: {item.get('evidence', 'Geen bewijs.')}", "• Verbeteradviezen:"])
                sug1 = item.get("suggestion_1") or item.get("suggestie_1") or item.get("advies_1")
                sug2 = item.get("suggestion_2") or item.get("suggestie_2") or item.get("advies_2")
                sug3 = item.get("suggestion_3") or item.get("suggestie_3") or item.get("advies_3")
                if sug1: lines.append(f"  - Suggestie 1: {sug1}")
                if sug2: lines.append(f"  - Suggestie 2: {sug2}")
                if sug3: lines.append(f"  - Suggestie 3: {sug3}")
    return "\n".join(lines)

def build_lesson_text(lesson, phases) -> str:
    return export_lesson_to_text(lesson, phases)

def export_lesson_to_pptx(lesson, phases, student_goals=None) -> io.BytesIO:
    if not PPTX_AVAILABLE: return None
    prs = Presentation(); prs.slide_width = PptInches(13.333); prs.slide_height = PptInches(7.5); blank_layout = prs.slide_layouts[6]
    def add_background_card(slide, left, top, width, height, bg_color=PptRGBColor(248, 250, 252), border_color=PptRGBColor(226, 232, 240)):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid(); shape.fill.fore_color.rgb = bg_color; shape.line.color.rgb = border_color; shape.line.width = PptPt(1)
        return shape
    slide1 = prs.slides.add_slide(blank_layout)
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg1.fill.solid(); bg1.fill.fore_color.rgb = PptRGBColor(30, 58, 138); bg1.line.fill.background()
    tb1 = slide1.shapes.add_textbox(PptInches(1.2), PptInches(2.2), PptInches(11), PptInches(3))
    tf1 = tb1.text_frame; tf1.word_wrap = True
    p_badge = tf1.paragraphs[0]; p_badge.text = "WIJZE LESSEN • LESPRESENTATIE"; p_badge.font.size = PptPt(12); p_badge.font.bold = True; p_badge.font.color.rgb = PptRGBColor(147, 197, 253); p_badge.space_after = PptPt(14)
    p_title = tf1.add_paragraph(); p_title.text = lesson.get('title') or 'Lesvoorbereiding'; p_title.font.size = PptPt(36); p_title.font.bold = True; p_title.font.color.rgb = PptRGBColor(255, 255, 255); p_title.space_after = PptPt(16)
    p_meta = tf1.add_paragraph(); p_meta.text = f"Vak: {lesson.get('subject', 'Nederlands')}  |  Klas: {lesson.get('class_name', '—')}  |  Datum: {lesson.get('date', '—')}  |  Duur: {lesson.get('duration', 50)} min."; p_meta.font.size = PptPt(14); p_meta.font.color.rgb = PptRGBColor(229, 231, 235)
    slide2 = prs.slides.add_slide(blank_layout)
    add_background_card(slide2, PptInches(0.8), PptInches(0.8), PptInches(11.733), PptInches(5.9))
    tb2 = slide2.shapes.add_textbox(PptInches(1.2), PptInches(1.1), PptInches(11), PptInches(5.3))
    tf2 = tb2.text_frame; tf2.word_wrap = True
    p2_top = tf2.paragraphs[0]; p2_top.text = "🎯 WAT LEER JE VANDAAG?"; p2_top.font.size = PptPt(24); p2_top.font.bold = True; p2_top.font.color.rgb = PptRGBColor(30, 58, 138); p2_top.space_after = PptPt(4)
    p2_sub = tf2.add_paragraph(); p2_sub.text = "Succescriteria & Doelen (Bouwsteen 2)"; p2_sub.font.size = PptPt(12); p2_sub.font.italic = True; p2_sub.font.color.rgb = PptRGBColor(100, 116, 139); p2_sub.space_after = PptPt(20)
    goals_to_show = student_goals if student_goals else [g.strip() for g in (lesson.get('learning_objectives') or '').split('\n') if g.strip()]
    if not goals_to_show: goals_to_show = ["De leerlingen bereiken de geplande doelstellingen van deze les."]
    for g_text in goals_to_show[:5]:
        p_g = tf2.add_paragraph(); clean_g = g_text.lstrip('•-123456789. '); p_g.text = f"✔  {clean_g}"; p_g.font.size = PptPt(17); p_g.font.color.rgb = PptRGBColor(30, 41, 59); p_g.space_after = PptPt(12)
    notes_slide2 = slide2.notes_slide.notes_text_frame; notes_slide2.text = f"FORMELE LESDOELEN (LERAAR):\n{lesson.get('learning_objectives') or 'Geen'}\n\nGEKOPPELDE LEERPLANDOELEN:\n{', '.join(normalize_multi_value(lesson.get('curriculum_goals')))}"
    normalized_phases = normalize_phase_list([dict(p) if hasattr(p, "keys") else p for p in phases], default_middle_count=1)
    for p_raw in normalized_phases:
        phase = normalize_phase(p_raw); slide_p = prs.slides.add_slide(blank_layout)
        add_background_card(slide_p, PptInches(0.8), PptInches(0.8), PptInches(11.733), PptInches(5.9))
        tb_p = slide_p.shapes.add_textbox(PptInches(1.2), PptInches(1.0), PptInches(11), PptInches(5.4))
        tf_p = tb_p.text_frame; tf_p.word_wrap = True
        p_p_title = tf_p.paragraphs[0]; p_p_title.text = phase['title'].upper(); p_p_title.font.size = PptPt(22); p_p_title.font.bold = True; p_p_title.font.color.rgb = PptRGBColor(30, 58, 138); p_p_title.space_after = PptPt(4)
        p_badge_bar = tf_p.add_paragraph()
        badge_elements = [f"⏱️ {phase['duration']} min."]
        if phase['activity_type']: badge_elements.append(f"👥 {format_multi_value(phase['activity_type'])}")
        if phase['materials']: badge_elements.append(f"📖 {format_multi_value(phase['materials'])}")
        p_badge_bar.text = "   •   ".join(badge_elements); p_badge_bar.font.size = PptPt(12); p_badge_bar.font.bold = True; p_badge_bar.font.color.rgb = PptRGBColor(37, 99, 235); p_badge_bar.space_after = PptPt(16)
        if phase['learning_content']:
            p_cont = tf_p.add_paragraph(); p_cont.text = f"Kernbegrippen: {phase['learning_content']}"; p_cont.font.size = PptPt(14); p_cont.font.bold = True; p_cont.font.color.rgb = PptRGBColor(71, 85, 105); p_cont.space_after = PptPt(10)
        if phase['organization']:
            p_org_title = tf_p.add_paragraph(); p_org_title.text = "Opdracht & Verloop:"; p_org_title.font.size = PptPt(15); p_org_title.font.bold = True; p_org_title.font.color.rgb = PptRGBColor(15, 23, 42); p_org_title.space_after = PptPt(4)
            for step in phase['organization'].split('\n')[:4]:
                if step.strip(): p_step = tf_p.add_paragraph(); p_step.text = f"• {step.strip()}"; p_step.font.size = PptPt(14); p_step.font.color.rgb = PptRGBColor(51, 65, 85); p_step.space_after = PptPt(4)
        if phase['student_activity']:
            p_act = tf_p.add_paragraph(); p_act.text = f"Jouw actie: {phase['student_activity']}"; p_act.font.size = PptPt(13); p_act.font.italic = True; p_act.font.color.rgb = PptRGBColor(30, 58, 138); p_act.space_before = PptPt(8)
        notes_p = slide_p.notes_slide.notes_text_frame; note_lines = [f"Fase: {phase['title']} ({phase['duration']} min.)"]
        if phase['teacher_activity']: note_lines.append(f"\nINSTRUCTIE / ACTIE LEERKRACHT:\n{phase['teacher_activity']}")
        if phase['lesson_objectives']: note_lines.append(f"\nLESDOEL VAN DEZE FASE:\n{phase['lesson_objectives']}")
        notes_p.text = "\n".join(note_lines)
    bio = io.BytesIO(); prs.save(bio); bio.seek(0); return bio

# ============================================================
# DATABASE LAAG
# ============================================================

def get_connection():
    connection = sqlite3.connect(DB_FILE, timeout=10, check_same_thread=False); connection.row_factory = sqlite3.Row; return connection

def init_database():
    LEERPLANNEN_DIR.mkdir(parents=True, exist_ok=True); DRAFT_FILE.parent.mkdir(parents=True, exist_ok=True)
    connection = get_connection(); cursor = connection.cursor()
    cursor.execute("""CREATE TABLE IF NOT EXISTS lessons (id INTEGER PRIMARY KEY AUTOINCREMENT, title TEXT NOT NULL, class_name TEXT, grade TEXT, study_direction TEXT, date TEXT, duration INTEGER, topic TEXT, learning_objectives TEXT, notes TEXT, analysis_json TEXT, created_at TEXT NOT NULL)""")
    cursor.execute("""CREATE TABLE IF NOT EXISTS lesson_phases (id INTEGER PRIMARY KEY AUTOINCREMENT, lesson_id INTEGER NOT NULL, phase_order INTEGER NOT NULL, title TEXT, duration INTEGER, teacher_activity TEXT, student_activity TEXT, activity_type TEXT, materials TEXT, FOREIGN KEY (lesson_id) REFERENCES lessons(id))""")
    cursor.execute("""CREATE TABLE IF NOT EXISTS lesson_versions (id INTEGER PRIMARY KEY AUTOINCREMENT, lesson_id INTEGER NOT NULL, version_number INTEGER NOT NULL, title TEXT, snapshot_json TEXT NOT NULL, created_at TEXT NOT NULL, FOREIGN KEY (lesson_id) REFERENCES lessons(id))""")
    cursor.execute("""CREATE TABLE IF NOT EXISTS lesson_progress (id INTEGER PRIMARY KEY AUTOINCREMENT, lesson_id INTEGER NOT NULL, eindpunt TEXT NOT NULL, timestamp TEXT NOT NULL, FOREIGN KEY (lesson_id) REFERENCES lessons(id))""")
    cursor.execute("PRAGMA table_info(lessons)"); lesson_columns = {column[1] for column in cursor.fetchall()}
    if "analysis_json" not in lesson_columns: cursor.execute("ALTER TABLE lessons ADD COLUMN analysis_json TEXT")
    if "curriculum_goals" not in lesson_columns: cursor.execute("ALTER TABLE lessons ADD COLUMN curriculum_goals TEXT")
    if "network" not in lesson_columns: cursor.execute("ALTER TABLE lessons ADD COLUMN network TEXT")
    if "subject" not in lesson_columns: cursor.execute("ALTER TABLE lessons ADD COLUMN subject TEXT")
    cursor.execute("PRAGMA table_info(lesson_phases)"); phase_columns = {column[1] for column in cursor.fetchall()}
    migration_columns = {"lesson_objectives": "TEXT", "learning_content": "TEXT", "organization": "TEXT", "section": "TEXT", "phase_number": "INTEGER"}
    for column_name, column_type in migration_columns.items():
        if column_name not in phase_columns: cursor.execute(f"ALTER TABLE lesson_phases ADD COLUMN {column_name} {column_type}")
    connection.commit(); connection.close()

def save_lesson_progress(lesson_id, eindpunt):
    connection = get_connection(); cursor = connection.cursor(); timestamp = datetime.now().isoformat()
    cursor.execute("INSERT INTO lesson_progress (lesson_id, eindpunt, timestamp) VALUES (?, ?, ?)", (lesson_id, eindpunt, timestamp)); connection.commit(); connection.close()

def get_lesson_progress(lesson_id):
    connection = get_connection(); cursor = connection.cursor()
    cursor.execute("SELECT id, eindpunt, timestamp FROM lesson_progress WHERE lesson_id = ? ORDER BY timestamp DESC", (lesson_id,))
    results = cursor.fetchall(); connection.close(); return [{"id": r["id"], "eindpunt": r["eindpunt"], "timestamp": r["timestamp"]} for r in results]

def delete_lesson_progress(progress_id):
    connection = get_connection(); cursor = connection.cursor()
    cursor.execute("DELETE FROM lesson_progress WHERE id = ?", (progress_id,)); connection.commit(); connection.close()

def insert_phases(cursor, lesson_id, phases):
    phases = renumber_middle_phases(phases)
    for order, raw_phase in enumerate(phases, start=1):
        phase = normalize_phase(raw_phase)
        cursor.execute("""INSERT INTO lesson_phases (lesson_id, phase_order, title, duration, teacher_activity, student_activity, activity_type, materials, lesson_objectives, learning_content, organization, section, phase_number) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""", (lesson_id, order, phase["title"], phase["duration"], phase["teacher_activity"], phase["student_activity"], serialize_multi_value(phase["activity_type"]), serialize_multi_value(phase["materials"]), phase["lesson_objectives"], phase["learning_content"], phase["organization"], phase["section"], phase["phase_number"]))

def save_lesson(title, class_name, grade, study_direction, lesson_date, duration, topic, learning_objectives, notes, phases, network="GO!", subject="Nederlands", curriculum_goals=None):
    phases = renumber_middle_phases(phases); connection = get_connection(); cursor = connection.cursor()
    cursor.execute("""INSERT INTO lessons (title, class_name, grade, study_direction, date, duration, topic, learning_objectives, notes, network, subject, curriculum_goals, analysis_json, created_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""", (title, class_name, grade, study_direction, lesson_date, duration, topic, learning_objectives, notes, network, subject, serialize_multi_value(curriculum_goals), None, datetime.now().isoformat()))
    lesson_id = cursor.lastrowid; insert_phases(cursor, lesson_id, phases); connection.commit(); connection.close(); return lesson_id

def update_lesson(lesson_id, title, class_name, grade, study_direction, lesson_date, duration, topic, learning_objectives, notes, phases, network="GO!", subject="Nederlands", curriculum_goals=None):
    phases = renumber_middle_phases(phases); connection = get_connection(); cursor = connection.cursor()
    old_lesson = connection.execute("SELECT * FROM lessons WHERE id = ?", (lesson_id,)).fetchone()
    old_phases = connection.execute("SELECT * FROM lesson_phases WHERE lesson_id = ? ORDER BY phase_order", (lesson_id,)).fetchall()
    if old_lesson:
        old_lesson_dict = dict(old_lesson); cursor.execute("SELECT COUNT(*) FROM lesson_versions WHERE lesson_id = ?", (lesson_id,)); new_ver_num = cursor.fetchone()[0] + 1
        snapshot = {"title": old_lesson_dict["title"], "class_name": old_lesson_dict["class_name"], "grade": old_lesson_dict["grade"], "study_direction": old_lesson_dict["study_direction"], "date": old_lesson_dict["date"], "duration": old_lesson_dict["duration"], "topic": old_lesson_dict["topic"], "learning_objectives": old_lesson_dict["learning_objectives"], "notes": old_lesson_dict["notes"], "network": old_lesson_dict["network"], "subject": old_lesson_dict["subject"], "curriculum_goals": normalize_multi_value(old_lesson_dict["curriculum_goals"]), "analysis_json": old_lesson_dict["analysis_json"], "phases": [dict(p) for p in old_phases]}
        cursor.execute("INSERT INTO lesson_versions (lesson_id, version_number, title, snapshot_json, created_at) VALUES (?, ?, ?, ?, ?)", (lesson_id, new_ver_num, old_lesson_dict["title"], json.dumps(snapshot, ensure_ascii=False), datetime.now().isoformat()))
        cursor.execute("""UPDATE lessons SET title = ?, class_name = ?, grade = ?, study_direction = ?, date = ?, duration = ?, topic = ?, learning_objectives = ?, notes = ?, network = ?, subject = ?, curriculum_goals = ? WHERE id = ?""", (title, class_name, grade, study_direction, lesson_date, duration, topic, learning_objectives, notes, network, subject, serialize_multi_value(curriculum_goals), lesson_id))
        cursor.execute("DELETE FROM lesson_phases WHERE lesson_id = ?", (lesson_id,)); insert_phases(cursor, lesson_id, phases); connection.commit(); connection.close()

def restore_lesson_version(lesson_id, version_id):
    connection = get_connection(); cursor = connection.cursor()
    cursor.execute("SELECT * FROM lesson_versions WHERE id = ?", (version_id,)); ver_row = cursor.fetchone()
    if not ver_row: connection.close(); return
    snapshot = json.loads(ver_row["snapshot_json"]); old_lesson = connection.execute("SELECT * FROM lessons WHERE id = ?", (lesson_id,)).fetchone()
    old_phases = connection.execute("SELECT * FROM lesson_phases WHERE lesson_id = ? ORDER BY phase_order", (lesson_id,)).fetchall()
    if old_lesson:
        old_lesson_dict = dict(old_lesson); cursor.execute("SELECT COUNT(*) FROM lesson_versions WHERE lesson_id = ?", (lesson_id,)); new_ver_num = cursor.fetchone()[0] + 1
        curr_snapshot = {"title": old_lesson_dict["title"], "class_name": old_lesson_dict["class_name"], "grade": old_lesson_dict["grade"], "study_direction": old_lesson_dict["study_direction"], "date": old_lesson_dict["date"], "duration": old_lesson_dict["duration"], "topic": old_lesson_dict["topic"], "learning_objectives": old_lesson_dict["learning_objectives"], "notes": old_lesson_dict["notes"], "network": old_lesson_dict["network"], "subject": old_lesson_dict["subject"], "curriculum_goals": normalize_multi_value(old_lesson_dict["curriculum_goals"]), "analysis_json": old_lesson_dict["analysis_json"], "phases": [dict(p) for p in old_phases]}
        cursor.execute("INSERT INTO lesson_versions (lesson_id, version_number, title, snapshot_json, created_at) VALUES (?, ?, ?, ?, ?)", (lesson_id, new_ver_num, old_lesson_dict["title"], json.dumps(curr_snapshot, ensure_ascii=False), datetime.now().isoformat()))
        cursor.execute("""UPDATE lessons SET title = ?, class_name = ?, grade = ?, study_direction = ?, date = ?, duration = ?, topic = ?, learning_objectives = ?, notes = ?, network = ?, subject = ?, curriculum_goals = ?, analysis_json = ? WHERE id = ?""", (snapshot.get("title"), snapshot.get("class_name"), snapshot.get("grade"), snapshot.get("study_direction"), snapshot.get("date"), snapshot.get("duration"), snapshot.get("topic"), snapshot.get("learning_objectives"), snapshot.get("notes"), snapshot.get("network"), snapshot.get("subject"), serialize_multi_value(snapshot.get("curriculum_goals")), snapshot.get("analysis_json"), lesson_id))
        cursor.execute("DELETE FROM lesson_phases WHERE lesson_id = ?", (lesson_id,)); insert_phases(cursor, lesson_id, snapshot.get("phases", [])); connection.commit(); connection.close()

def get_lesson_versions(lesson_id): connection = get_connection(); versions = connection.execute("SELECT * FROM lesson_versions WHERE lesson_id = ? ORDER BY version_number DESC", (lesson_id,)).fetchall(); connection.close(); return [dict(v) for v in versions]
def delete_lesson(lesson_id): connection = get_connection(); cursor = connection.cursor(); cursor.execute("DELETE FROM lesson_versions WHERE lesson_id = ?", (lesson_id,)); cursor.execute("DELETE FROM lesson_phases WHERE lesson_id = ?", (lesson_id,)); cursor.execute("DELETE FROM lesson_progress WHERE lesson_id = ?", (lesson_id,)); cursor.execute("DELETE FROM lessons WHERE id = ?", (lesson_id,)); connection.commit(); connection.close()
def get_lessons(): connection = get_connection(); lessons = connection.execute("SELECT * FROM lessons ORDER BY date DESC, id DESC").fetchall(); connection.close(); return [{k: l[k] for k in l.keys()} for l in lessons]
def get_lesson(lesson_id): connection = get_connection(); lesson = connection.execute("SELECT * FROM lessons WHERE id = ?", (lesson_id,)).fetchone(); phases = connection.execute("SELECT * FROM lesson_phases WHERE lesson_id = ? ORDER BY phase_order", (lesson_id,)).fetchall(); connection.close(); lesson_dict = {k: lesson[k] for k in lesson.keys()} if lesson else {}; phases_list = [{k: p[k] for k in p.keys()} for p in phases]; return lesson_dict, phases_list

def db_phases_to_editor(phases):
    raw = [dict(phase) for phase in phases]
    if not raw: return normalize_phase_list([], default_middle_count=3)
    has_sections = any(item.get("section") in VALID_SECTIONS for item in raw)
    if not has_sections:
        converted = []
        if len(raw) == 1:
            existing = raw[0]; existing["section"], existing["phase_number"] = "LESMIDDEN", 1; converted = [empty_phase("LESINTRO"), existing, empty_phase("LESAFSLUITING")]
        else:
            for index, phase in enumerate(raw):
                if index == 0: phase["section"], phase["phase_number"] = "LESINTRO", None
                elif index == len(raw) - 1: phase["section"], phase["phase_number"] = "LESAFSLUITING", None
                else: phase["section"], phase["phase_number"] = "LESMIDDEN", index
                converted.append(phase)
        raw = converted
    for phase in raw:
        if not phase.get("phase_id"): phase["phase_id"] = f"db_{phase['id']}" if phase.get("id") is not None else create_phase_id()
    return normalize_phase_list(raw, default_middle_count=1)

def save_ai_analysis(lesson_id, analysis):
    if isinstance(analysis, dict):
        if "analyzed_at" not in analysis: analysis["analyzed_at"] = datetime.now().strftime("%d-%m-%Y om %H:%M")
        connection = get_connection(); connection.execute("UPDATE lessons SET analysis_json = ? WHERE id = ?", (json.dumps(analysis, ensure_ascii=False), lesson_id)); connection.commit(); connection.close()

def get_ai_analysis(lesson):
    analysis_json = lesson.get("analysis_json")
    if not analysis_json: return None
    try: return json.loads(analysis_json)
    except (json.JSONDecodeError, TypeError): return None

# ============================================================
# CSS & WEERGAVE
# ============================================================

def apply_custom_css():
    st.markdown(
        """
        <style>
        html { scroll-behavior: smooth; }
        .lesson-help { color: #666; font-size: 0.85rem; }
        
        .detail-badge {
            display: inline-flex !important;
            align-items: center !important;
            background: rgba(59, 130, 246, 0.12) !important;
            border: 1px solid rgba(59, 130, 246, 0.28) !important;
            color: #93c5fd !important;
            padding: 5px 12px !important;
            border-radius: 8px !important;
            font-size: 0.86rem !important;
            font-weight: 500 !important;
        }

        button[data-testid="stBaseButton-secondary"] {
            border: 1px solid rgba(255, 255, 255, 0.16) !important;
            border-radius: 8px !important;
            transition: all 0.2s ease !important;
            background: rgba(255, 255, 255, 0.03) !important;
            color: #f3f4f6 !important;
        }

        button[data-testid="stBaseButton-secondary"]:hover {
            border-color: #3b82f6 !important;
            color: #ffffff !important;
            background: rgba(59, 130, 246, 0.22) !important;
            box-shadow: 0 4px 12px rgba(59, 130, 246, 0.25) !important;
            transform: translateY(-1px) !important;
        }

        /* Zorg dat text areas automatisch meegroeien met inhoud */
        .stTextArea textarea {
            field-sizing: content;
        }

        .stTextArea textarea[aria-label="Ruwe lesvoorbereiding / Brontekst"] { min-height: 200px !important; }
        .stTextArea textarea[aria-label="Algemene lesdoelen"], 
        .stTextArea textarea[aria-label="Algemene leerdoelen"] { min-height: 100px !important; }
        .stTextArea textarea[aria-label="Algemene opmerkingen"] { min-height: 80px !important; }
        .stTextArea textarea[aria-label="Lesdoelen"], 
        .stTextArea textarea[aria-label="Leerinhouden"] { min-height: 120px !important; }
        
        /* Lesverloop (Organisatie) start dubbel zo groot bij leeg veld */
        .stTextArea textarea[aria-label="Organisatie"] { min-height: 240px !important; }

        details.clean-card {
            position: relative !important;
            background: rgba(255, 255, 255, 0.04) !important;
            border: 1px solid rgba(255, 255, 255, 0.12) !important;
            border-radius: 10px !important;
            padding: 12px 14px !important;
            margin-bottom: 14px !important;
            color: #ffffff !important;
            transition: all 0.2s ease-in-out !important;
            box-sizing: border-box !important;
            cursor: pointer !important;
        }
        
        details.clean-card:hover {
            background: rgba(59, 130, 246, 0.12) !important;
            border-color: #3b82f6 !important;
            box-shadow: 0 4px 14px rgba(0, 0, 0, 0.35) !important;
        }

        details.clean-card[open] {
            background: rgba(30, 58, 138, 0.22) !important;
            border-color: #3b82f6 !important;
            box-shadow: 0 6px 18px rgba(0, 0, 0, 0.45) !important;
        }
        
        details.clean-card summary {
            list-style: none !important;
            outline: none !important;
            user-select: none !important;
            display: flex !important;
            flex-direction: column !important;
        }
        
        details.clean-card summary::-webkit-details-marker {
            display: none !important;
        }

        .clean-card-header {
            display: flex !important;
            justify-content: space-between !important;
            align-items: center !important;
            margin-bottom: 6px !important;
        }
        
        .clean-card-title {
            font-weight: 600 !important;
            font-size: 0.95rem !important;
            color: #ffffff !important;
            line-height: 1.3 !important;
        }

        .clean-card-toggle {
            font-size: 0.75rem !important;
            color: #94a3b8 !important;
            transition: transform 0.2s ease !important;
        }

        details.clean-card[open] .clean-card-toggle {
            transform: rotate(180deg) !important;
            color: #60a5fa !important;
        }
        
        .clean-card-stars {
            display: flex !important;
            align-items: center !important;
            font-size: 0.9rem !important;
            letter-spacing: 1px !important;
            color: #fbbf24 !important;
        }
        
        .clean-card-score {
            font-size: 0.8rem !important;
            color: #9ca3af !important;
            margin-left: 6px !important;
            letter-spacing: normal !important;
        }
        
        .clean-card-suggestion {
            display: -webkit-box !important;
            -webkit-line-clamp: 2 !important;
            -webkit-box-orient: vertical !important;
            overflow: hidden !important;
            text-overflow: ellipsis !important;
            font-size: 0.8rem !important;
            color: #9ca3af !important;
            margin-top: 8px !important;
            padding-top: 8px !important;
            border-top: 1px solid rgba(255, 255, 255, 0.1) !important;
            font-style: italic !important;
            line-height: 1.4 !important;
        }

        details.clean-card[open] .clean-card-suggestion {
            display: none !important;
        }

        .clean-card-body {
            margin-top: 12px !important;
            padding-top: 10px !important;
            border-top: 1px solid rgba(59, 130, 246, 0.35) !important;
            font-size: 0.85rem !important;
            line-height: 1.45 !important;
            color: #e2e8f0 !important;
            cursor: default !important;
        }

        .clean-card-fullname {
            font-size: 0.82rem !important;
            color: #93c5fd !important;
            margin-bottom: 8px !important;
            font-weight: 500 !important;
        }

        .clean-card-badge {
            display: inline-block !important;
            font-size: 0.72rem !important;
            padding: 2px 7px !important;
            border-radius: 4px !important;
            background: rgba(59, 130, 246, 0.2) !important;
            color: #93c5fd !important;
            border: 1px solid rgba(59, 130, 246, 0.4) !important;
            margin-bottom: 8px !important;
        }

        .clean-card-section-label {
            font-weight: 600 !important;
            color: #f1f5f9 !important;
            font-size: 0.8rem !important;
            margin-top: 8px !important;
            margin-bottom: 2px !important;
            display: flex !important;
            align-items: center !important;
            gap: 4px !important;
        }

        .clean-card-evidence {
            color: #cbd5e1 !important;
            font-size: 0.82rem !important;
            background: rgba(0, 0, 0, 0.2) !important;
            padding: 6px 8px !important;
            border-radius: 6px !important;
            border-left: 3px solid #3b82f6 !important;
            margin-bottom: 8px !important;
        }

        .clean-card-adviezen-list {
            margin: 4px 0 0 0 !important;
            padding-left: 16px !important;
            color: #f8fafc !important;
            font-size: 0.82rem !important;
        }

        .clean-card-adviezen-list li {
            margin-bottom: 6px !important;
        }
        
        .tracker-card { border-radius: 8px; padding: 12px 16px; margin-bottom: 10px; border: 1px solid rgba(255, 255, 255, 0.08); transition: all 0.2s ease; }
        .tracker-card-covered { background-color: rgba(16, 185, 129, 0.08); border-left: 5px solid #10b981; }
        .tracker-card-uncovered { background-color: rgba(255, 255, 255, 0.02); border-left: 5px solid #6b7280; }
        </style>
        """, unsafe_allow_html=True
    )

def show_ai_analysis(analysis, lesson_obj=None, version_str="Actuele versie"):
    if not analysis:
        st.error("De AI-analyzer gaf geen resultaat terug.")
        return

    st.divider()
    st.subheader("🤖 AI-didactische analyse")

    analyzed_time = analysis.get("analyzed_at", "Niet geregistreerd")
    st.caption(f"🕒 **Geanalyseerd op:** {analyzed_time}  |  🏷️ **{version_str}**")

    if lesson_obj:
        if PDF_AVAILABLE:
            pdf_buf = export_analysis_to_pdf(lesson_obj, analysis, version_str=version_str)
            if pdf_buf:
                clean_title = "".join(c for c in (lesson_obj.get('title') or "les") if c.isalnum() or c in (' ', '_', '-')).strip().replace(' ', '_')
                st.download_button(
                    label="📥 Download AI-Analyse (PDF)",
                    data=pdf_buf,
                    file_name=f"AI_Analyse_{clean_title}_{lesson_obj.get('date') or 'export'}.pdf",
                    mime="application/pdf",
                    key=f"dl_pdf_analysis_{lesson_obj.get('id', 'temp')}_{analysis.get('analyzed_at', '')}",
                    type="secondary",
                    help="Exporteer deze didactische evaluatie als PDF-document"
                )
        else:
            st.caption("ℹ️ *Installeer `reportlab` om PDF-exports in te schakelen: `pip install reportlab`*")

    total_score = analysis.get("total_score")
    if total_score is not None: 
        st.metric("Totale score", f"{total_score} / 60")
    if analysis.get("summary"): 
        st.info(analysis["summary"])

    principles = analysis.get("principles", [])
    if principles:
        principles_sorted = sorted(principles, key=get_principle_number)
        st.markdown("### 📋 Didactische Rapportkaart")
        for i in range(0, len(principles_sorted), 3):
            cols = st.columns(3)
            row_items = principles_sorted[i:i+3]
            for col, item in zip(cols, row_items):
                p_obj = get_principle(item)
                stars, score_num = format_star_rating(item.get("score", 0))
                sug1 = item.get("suggestion_1") or item.get("suggestie_1") or item.get("advies_1")
                sug2 = item.get("suggestion_2") or item.get("suggestie_2") or item.get("advies_2")
                sug3 = item.get("suggestion_3") or item.get("suggestie_3") or item.get("advies_3")
                first_sug = sug1 or sug2 or sug3 or "Geen specifiek verbeteradvies."
                evidence = item.get("evidence") or item.get("bewijs") or "Geen specifiek bewijs gevonden in het lesplan."
                status = item.get("status", "")

                card_html = f"""
                <details class="clean-card">
                    <summary>
                        <div class="clean-card-header">
                            <span class="clean-card-title">{p_obj['emoji']} B{p_obj['number']}. {p_obj['short']}</span>
                            <span class="clean-card-toggle">▼</span>
                        </div>
                        <div class="clean-card-stars">{stars} <span class="clean-card-score">({score_num}/5)</span></div>
                        <div class="clean-card-suggestion">💡 {html.escape(str(first_sug))}</div>
                    </summary>
                    <div class="clean-card-body">
                        <div class="clean-card-fullname">{html.escape(p_obj['name'])}</div>
                        <div class="clean-card-section-label">🔍 Bewijs uit de les:</div>
                        <div class="clean-card-evidence">{html.escape(str(evidence))}</div>
                    </div>
                </details>
                """
                with col:
                    st.markdown(card_html, unsafe_allow_html=True)

# ============================================================
# FASE-EDITOR WIDGETS (MET INTELLIGENTE '+' KNOP PER FASE)
# ============================================================

def render_phase_row(stored, row_index, prefix, lesson_id=None, allow_delete=False):
    phase = normalize_phase(stored)
    phase_id = phase["phase_id"]
    columns = st.columns([1.15, 2.0, 2.0, 2.6, 0.8])

    if phase["section"] == "LESINTRO": label = "LESINTRO"
    elif phase["section"] == "LESAFSLUITING": label = "LESAFSLUITING"
    else: label = f"Lesfase {phase['phase_number'] or row_index}"

    with columns[0]:
        if phase["section"] == "LESMIDDEN":
            st.markdown("**LESMIDDEN**")
            st.caption(label)
            btn_col1, btn_col2 = st.columns(2)
            with btn_col1:
                if allow_delete:
                    if st.button("🗑️", key=f"{prefix}_delete_{lesson_id or 'new'}_{phase_id}", help="Verwijder deze lesfase"):
                        return {"_delete": True}
            with btn_col2:
                if st.button("➕", key=f"{prefix}_add_after_{lesson_id or 'new'}_{phase_id}", help="Voeg hierna een nieuwe lesfase in"):
                    return {"_add_after": row_index}
        else: 
            st.markdown(f"**{label}**")
            if st.button("➕", key=f"{prefix}_add_after_{lesson_id or 'new'}_{phase_id}", help="Voeg hierna een nieuwe lesfase in"):
                return {"_add_after": row_index}

    base = f"{prefix}_{lesson_id or 'new'}_{phase_id}"
    with columns[1]: objectives = st.text_area("Lesdoelen", value=phase["lesson_objectives"], key=f"{base}_obj", label_visibility="collapsed", placeholder="Wat moeten leerlingen bereiken?")
    with columns[2]: content = st.text_area("Leerinhouden", value=phase["learning_content"], key=f"{base}_content", label_visibility="collapsed", placeholder="Welke leerstof komt aan bod?")
    with columns[3]:
        valid_act = [x for x in phase["activity_type"] if x in ACTIVITY_OPTIONS]
        valid_mat = [x for x in phase["materials"] if x in MATERIAL_OPTIONS]
        activities = st.multiselect("Werkvorm", options=ACTIVITY_OPTIONS, default=valid_act, key=f"{base}_activity", placeholder="Werkvorm", label_visibility="collapsed")
        materials = st.multiselect("Media", options=MATERIAL_OPTIONS, default=valid_mat, key=f"{base}_materials", placeholder="Media", label_visibility="collapsed")
        organization = st.text_area("Organisatie", value=phase["organization"], key=f"{base}_org", placeholder="Lesverloop...", label_visibility="collapsed")
    with columns[4]: duration = st.number_input("Tijd", min_value=1, max_value=120, value=phase["duration"], step=1, key=f"{base}_time", label_visibility="collapsed")

    return {**phase, "title": label, "lesson_objectives": objectives, "learning_content": content, "organization": organization, "duration": duration, "activity_type": activities, "materials": materials}

def render_new_lesson_structure():
    phases = normalize_phase_list(st.session_state.get("new_phases", []), default_middle_count=3)
    st.markdown("### Lesvoorbereiding")
    st.info(f"Deze les heeft momenteel **{len(phases) - 2} lesfase(n)** in het midden.")

    header = st.columns([1.15, 2.0, 2.0, 2.6, 0.8])
    for column, label in zip(header, ["Fase", "Lesdoelen", "Leerinhouden", "Werkvormen / media / organisatie", "Tijd"]):
        with column: st.markdown(f"**{label}**")

    edited, delete_phase_id, add_after_idx = [], None, None
    for row_index, phase in enumerate(phases):
        result = render_phase_row(phase, row_index, prefix="new", allow_delete=(phase["section"] == "LESMIDDEN"))
        if result.get("_delete"): 
            delete_phase_id = phase["phase_id"]
            edited.append(phase)
        elif result.get("_add_after") is not None:
            add_after_idx = result["_add_after"]
            edited.append(phase)
        else:
            edited.append(result)
        st.divider()

    # Intelligente toevoeging: voegt de nieuwe fase direct na de gekozen fase in
    # Bewaakt automatisch dat LESAFSLUITING altijd het sluitstuk blijft
    if add_after_idx is not None:
        intro = next((p for p in edited if p.get("section") == "LESINTRO"), empty_phase("LESINTRO"))
        closing = next((p for p in edited if p.get("section") == "LESAFSLUITING"), empty_phase("LESAFSLUITING"))
        middles = [p for p in edited if p.get("section") == "LESMIDDEN"]

        if add_after_idx == 0:
            insert_pos = 0
        elif add_after_idx > len(middles):
            insert_pos = len(middles)
        else:
            insert_pos = add_after_idx

        middles.insert(insert_pos, empty_phase("LESMIDDEN", insert_pos + 1))
        st.session_state["new_phases"] = renumber_middle_phases([intro, *middles, closing])
        st.rerun()

    if st.button("＋ Lesfase toevoegen aan het midden", key="new_add_phase"):
        intro = next((p for p in edited if p.get("section") == "LESINTRO"), empty_phase("LESINTRO"))
        closing = next((p for p in edited if p.get("section") == "LESAFSLUITING"), empty_phase("LESAFSLUITING"))
        middles = [p for p in edited if p.get("section") == "LESMIDDEN"]
        middles.append(empty_phase("LESMIDDEN", len(middles) + 1))
        st.session_state["new_phases"] = renumber_middle_phases([intro, *middles, closing])
        st.rerun()

    if delete_phase_id is not None:
        if len(phases) - 2 <= 1: 
            st.warning("Een les moet minstens één lesfase in het midden hebben.")
        else:
            st.session_state["new_phases"] = renumber_middle_phases([p for p in edited if p.get("phase_id") != delete_phase_id])
            st.rerun()

    result = renumber_middle_phases(edited)
    st.session_state["new_phases"] = result
    return result

def display_lesson_structure(phases):
    normalized = normalize_phase_list([dict(p) if hasattr(p, "keys") else p for p in phases], default_middle_count=1)
    st.markdown("### Lesvoorbereiding")
    header = st.columns([1.15, 2.0, 2.0, 2.6, 0.8])
    for column, label in zip(header, ["Fase", "Lesdoelen", "Leerinhouden", "Werkvormen / media / organisatie", "Tijd"]):
        with column: st.markdown(f"**{label}**")

    for phase in normalized:
        columns = st.columns([1.15, 2.0, 2.0, 2.6, 0.8])
        with columns[0]:
            if phase["section"] == "LESINTRO": st.markdown("**LESINTRO**")
            elif phase["section"] == "LESAFSLUITING": st.markdown("**LESAFSLUITING**")
            else: st.markdown("**LESMIDDEN**"); st.caption(f"Lesfase {phase['phase_number']}")
        with columns[1]: st.write(phase["lesson_objectives"] or "—")
        with columns[2]: st.write(phase["learning_content"] or "—")
        with columns[3]:
            st.write(f"**Werkvorm:** {format_multi_value(phase['activity_type'])}")
            st.write(f"**Media:** {format_multi_value(phase['materials'])}")
            if phase["organization"]: st.write(f"**Organisatie:** {phase['organization']}")
        with columns[4]: st.write(f"{phase['duration']} min.")
        st.divider()

def render_edit_structure(existing, lesson_id):
    phases = normalize_phase_list(existing, default_middle_count=1)
    st.info(f"Deze les heeft momenteel **{len(phases) - 2} lesfase(n)** in het midden.")

    header = st.columns([1.15, 2.0, 2.0, 2.6, 0.8])
    for column, label in zip(header, ["Fase", "Lesdoelen", "Leerinhouden", "Werkvormen / media / organisatie", "Tijd"]):
        with column: st.markdown(f"**{label}**")

    edited, delete_phase_id, add_after_idx = [], None, None
    for row_index, phase in enumerate(phases):
        result = render_phase_row(phase, row_index, prefix="edit", lesson_id=lesson_id, allow_delete=(phase["section"] == "LESMIDDEN"))
        if result.get("_delete"): 
            delete_phase_id = phase["phase_id"]
            edited.append(phase)
        elif result.get("_add_after") is not None:
            add_after_idx = result["_add_after"]
            edited.append(phase)
        else:
            edited.append(result)
        st.divider()

    if add_after_idx is not None:
        intro = next((p for p in edited if p.get("section") == "LESINTRO"), empty_phase("LESINTRO"))
        closing = next((p for p in edited if p.get("section") == "LESAFSLUITING"), empty_phase("LESAFSLUITING"))
        middles = [p for p in edited if p.get("section") == "LESMIDDEN"]

        if add_after_idx == 0:
            insert_pos = 0
        elif add_after_idx > len(middles):
            insert_pos = len(middles)
        else:
            insert_pos = add_after_idx

        middles.insert(insert_pos, empty_phase("LESMIDDEN", insert_pos + 1))
        st.session_state[f"edit_phases_{lesson_id}"] = renumber_middle_phases([intro, *middles, closing])
        st.rerun()

    if st.button("＋ Lesfase toevoegen aan het midden", key=f"edit_add_phase_{lesson_id}"):
        intro = next((p for p in edited if p.get("section") == "LESINTRO"), empty_phase("LESINTRO"))
        closing = next((p for p in edited if p.get("section") == "LESAFSLUITING"), empty_phase("LESAFSLUITING"))
        middles = [p for p in edited if p.get("section") == "LESMIDDEN"]
        middles.append(empty_phase("LESMIDDEN", len(middles) + 1))
        st.session_state[f"edit_phases_{lesson_id}"] = renumber_middle_phases([intro, *middles, closing])
        st.rerun()

    if delete_phase_id is not None:
        if len(phases) - 2 <= 1: 
            st.warning("Een les moet minstens één lesfase in het midden hebben.")
        else:
            st.session_state[f"edit_phases_{lesson_id}"] = renumber_middle_phases([p for p in edited if p.get("phase_id") != delete_phase_id])
            st.rerun()

    result = renumber_middle_phases(edited)
    st.session_state[f"edit_phases_{lesson_id}"] = result
    return result

# ============================================================
# STREAMLIT CONFIGURATIE & INITIALISATIE
# ============================================================
st.set_page_config(page_title=APP_TITLE, page_icon="🧠", layout="wide")
init_database()
apply_custom_css()

# Initialiseer basis state voor provider en model indien nog niet gezet
if "ai_provider" not in st.session_state:
    st.session_state.ai_provider = "Google Gemini"

if "gekozen_model" not in st.session_state:
    try:
        st.session_state.gekozen_model = get_default_model(st.session_state.ai_provider)
    except Exception:
        st.session_state.gekozen_model = "gemini-3.6-flash"

st.sidebar.title("🧠 Wijze Lessen")
st.sidebar.markdown("---")
st.sidebar.subheader("🤖 AI-Instellingen")

# 1. Provider Selectie
provider_options = ["Google Gemini", "Lokaal (Ollama)"]
# Valideer en update sessie als er een onverwachte waarde instaat
if st.session_state.ai_provider not in provider_options:
    st.session_state.ai_provider = "Google Gemini"

provider_idx = provider_options.index(st.session_state.ai_provider)
selected_provider = st.sidebar.radio(
    "AI-provider", 
    options=provider_options, 
    index=provider_idx,
    key="provider_radio_widget"
)

# Detecteer wijziging en zet standaardwaarden terug
if selected_provider != st.session_state.ai_provider:
    st.session_state.ai_provider = selected_provider
    try:
        st.session_state.gekozen_model = get_default_model(selected_provider)
    except Exception:
        st.session_state.gekozen_model = "gemini-3.6-flash" if "Google" in selected_provider else "qwen2.5:7b"
    st.rerun()

# 2. Check Gemini beschikbaarheid en toon waarschuwing (geen crash!)
try:
    if st.session_state.ai_provider == "Google Gemini" and not google_available():
        st.sidebar.warning("⚠️ Google Gemini is niet beschikbaar. Controleer GEMINI_API_KEY in .streamlit/secrets.toml.")
except Exception:
    pass

# 3. Model Selectie
try:
    beschikbare_modellen = get_models(st.session_state.ai_provider)
except TypeError:
    beschikbare_modellen = get_models() # Fallback voor als provider param niet bestaat

if not beschikbare_modellen: 
    try:
        beschikbare_modellen = [get_default_model(st.session_state.ai_provider)]
    except Exception:
        beschikbare_modellen = ["gemini-3.6-flash"] if st.session_state.ai_provider == "Google Gemini" else ["gemma4:e4b"]

current_model = st.session_state.gekozen_model
if current_model not in beschikbare_modellen:
    current_model = beschikbare_modellen[0]
    st.session_state.gekozen_model = current_model

model_idx = beschikbare_modellen.index(current_model)
selected_model = st.sidebar.selectbox(
    "Kies AI-model:", 
    beschikbare_modellen, 
    index=model_idx,
    key="model_selectbox_widget"
)

if selected_model != st.session_state.gekozen_model:
    st.session_state.gekozen_model = selected_model
    st.rerun()

st.sidebar.markdown("---")

page = st.sidebar.radio("Navigatie", [
    "🏠 Dashboard", 
    "📝 Nieuwe les", 
    "📚 Mijn lessen", 
    "🎯 Doelen Tracker", 
    "🧱 De 12 bouwstenen"
])

# ============================================================
# PAGINA: 🏠 DASHBOARD
# ============================================================
if page == "🏠 Dashboard":
    st.title("🧠 Wijze Lessen Coach")
    st.caption("Een didactische coach gebaseerd op de 12 bouwstenen van Wijze Lessen.")
    lessons = get_lessons()
    col1, col2 = st.columns(2)
    with col1: st.metric("Opgeslagen lessen", len(lessons))
    with col2: st.metric("Didactische bouwstenen", 12)
    st.divider()
    st.subheader("📚 Recente lessen")
    if lessons:
        for lesson in lessons[:5]:
            with st.container(border=True):
                st.write(f"### {lesson.get('title')}")
                st.write(f"{lesson.get('class_name') or ''} · {lesson.get('grade') or ''} · {lesson.get('topic') or ''}")
                st.caption(f"Lesduur: {lesson.get('duration') or 0} minuten")
    else: st.info("Je hebt nog geen lessen opgeslagen.")

# ============================================================
# PAGINA: 📝 NIEUWE LES (MET BESTANDSUPLOAD & 2-STAPS AI)
# ============================================================
elif page == "📝 Nieuwe les":
    initialize_new_lesson_state()
    st.title("📝 Nieuwe les")
    st.write("Maak een lesvoorbereiding volgens het didactische sjabloon.")
    
    st.subheader("🎓 Leerplancontext")
    c_col1, c_col2, c_col3, c_col4 = st.columns(4)
    with c_col1: network = st.selectbox("Onderwijsnet", ["GO!", "Katholiek Onderwijs", "OVSG"], index=0, key="new_network")
    with c_col2: grade = st.selectbox("Graad / Leerjaar", ["1", "2", "3", "5", "6"], index=2, key="new_grade")
    with c_col3: study_direction = st.selectbox("Finaliteit", ["Dubbele finaliteit", "Doorstroom", "Arbeidsmarkt"], index=0, key="new_study_direction")
    with c_col4: subject = st.selectbox("Vak", ["Nederlands", "Wiskunde", "Geschiedenis", "Engels"], index=0, key="new_subject")

    available_goals = load_curriculum_goals(net=network, graad=grade, finaliteit=study_direction, vak=subject)
    if available_goals: st.caption(f"✅ **{len(available_goals)} officiële leerplandoelen** geladen voor {network} - {grade}e graad {study_direction} ({subject}).")
    else: st.caption("ℹ️ Geen specifiek leerplanbestand gevonden in `data/leerplannen/`. Handmatige invoer blijft mogelijk.")

    st.divider()
    st.subheader("✨ AI gebruiken (Document uploaden of tekst plakken)")
    st.caption("Upload een PDF, Word-document, PowerPoint-presentatie of plak je ruwe lesnotities. De AI doorloopt een 2-staps analyse: eerst een didactisch sterke lesopbouw ontwerpen, daarna automatisch structureren en leerplandoelen koppelen.")

    uploaded_doc = st.file_uploader(
        "📁 Upload een bronbestand (PDF, Word, PowerPoint of TXT)",
        type=["pdf", "docx", "pptx", "txt"],
        help="De AI leest alle tekst, dia's en sprekersnotities uit."
    )
    if uploaded_doc is not None:
        if st.session_state.get("_last_uploaded_file") != uploaded_doc.name:
            with st.spinner(f"Document '{uploaded_doc.name}' uitlezen..."):
                doc_text = extract_text_from_file(uploaded_doc)
                if doc_text.strip():
                    st.session_state["new_ai_text"] = doc_text
                    st.session_state["_last_uploaded_file"] = uploaded_doc.name
                    st.success(f"Bestand '{uploaded_doc.name}' succesvol uitgelezen en ingeladen in het tekstvak!")

    st.text_area("Ruwe lesvoorbereiding / Brontekst", key="new_ai_text", placeholder="Plak hier je lesnotities of upload hierboven een document...")

    ai_col1, ai_col2 = st.columns(2)
    with ai_col1: extract_clicked = st.button("✨ AI vult sjabloon in (2 stappen)", type="primary", use_container_width=True)
    with ai_col2: clear_clicked = st.button("🗑️ Formulier leegmaken", use_container_width=True)

    if clear_clicked:
        reset_new_lesson()
        st.rerun()

    if extract_clicked:
        if not st.session_state.get("new_ai_text", "").strip(): 
            st.warning("Plak eerst je ruwe lesvoorbereiding in het tekstvak of upload een bestand.")
        else:
            try:
                progress_bar = st.progress(10, text=f"Stap 1/2: AI ontwerpt didactische lesopbouw met '{st.session_state['gekozen_model']}'...")
                
                didactic_draft = create_didactic_draft(
                    st.session_state["new_ai_text"],
                    duration=int(st.session_state.get("new_duration", 50)),
                    provider=st.session_state.ai_provider,
                    model_name=st.session_state['gekozen_model']
                )

                progress_bar.progress(60, text="Stap 2/2: Structureren naar sjabloon en officiële leerplandoelen koppelen...")
                extracted = extract_lesson_from_text(
                    didactic_draft,
                    provider=st.session_state.ai_provider,
                    model_name=st.session_state['gekozen_model'],
                    available_goals=available_goals,
                    original_raw_text=st.session_state["new_ai_text"]
                )
                
                progress_bar.progress(90, text="Resultaten inladen in formulier...")
                apply_extracted_lesson_to_session(extracted, available_goals=available_goals)
                st.session_state["new_ai_extraction_done"] = True
                
                progress_bar.progress(100, text="✅ Les succesvol verwerkt!")
                st.rerun()
            except Exception as error:
                st.error("Er ging iets mis tijdens de AI-verwerking.")
                st.exception(error)

    if st.session_state.get("new_ai_extraction_done", False): 
        if st.session_state.get("new_ai_warning"):
            st.warning(st.session_state["new_ai_warning"])
        else:
            st.info("De gegevens hieronder zijn door AI ingevuld. Controleer ze altijd zelf.")
    st.divider()
    st.subheader("Basisinformatie")

    col1, col2 = st.columns(2)
    with col1:
        title = st.text_input("Lesnaam *", key="new_title", placeholder="Bijvoorbeeld: Reynaert de Vos — personages")
        class_name = st.text_input("Klas", key="new_class_name", placeholder="5DF")
    with col2:
        lesson_date = st.date_input("Datum", value=st.session_state.get("new_date", date.today()), key="new_date")
        duration = st.number_input("Totale lesduur (minuten)", min_value=1, max_value=240, value=int(st.session_state.get("new_duration", 50)), step=5, key="new_duration")

    topic = st.text_input("Onderwerp", key="new_topic")
    learning_objectives = st.text_area("Algemene lesdoelen", key="new_learning_objectives")

    curriculum_options = [f"{g['code']} - {g['omschrijving']} ({g.get('cluster', '')})" for g in available_goals]
    selected_curriculum_goals = st.multiselect("🎯 Gekoppelde Officiële Leerplandoelen (automatisch voorgesteld door AI)", options=curriculum_options, default=[item for item in st.session_state.get("new_curriculum_goals", []) if item in curriculum_options], key="selected_goals_widget")

    notes = st.text_area("Algemene opmerkingen", key="new_notes")
    st.divider()
    phases = render_new_lesson_structure()
    st.divider()
    st.markdown("### ⏱️ Tijdcontrole")
    show_duration_check(duration, phases)
    submitted = st.button("💾 Les opslaan", key="save_new_lesson", type="primary", use_container_width=True)

    if submitted:
        if not title.strip(): st.error("Geef eerst een naam voor de les.")
        elif not learning_objectives.strip(): st.error("Voeg minstens één algemeen leerdoel toe.")
        else:
            save_lesson(title=title, class_name=class_name, grade=grade, study_direction=study_direction, lesson_date=lesson_date.isoformat(), duration=duration, topic=topic, learning_objectives=learning_objectives, notes=notes, phases=phases, network=network, subject=subject, curriculum_goals=selected_curriculum_goals)
            st.success(f"Les '{title}' is succesvol opgeslagen.")
            st.session_state["new_ai_extraction_done"] = False
            reset_new_lesson()

# ============================================================
# PAGINA: 📚 MIJN LESSEN (MASTER-DETAIL MET INSCHUIVEND DETAILVENSTER)
# ============================================================
elif page == "📚 Mijn lessen":
    lessons = get_lessons()
    selected_id = st.session_state.get("selected_lesson_id")

    if selected_id is None:
        st.title("📚 Mijn lessen")
        st.caption("Klik op een les om de volledige fiche en didactische analyse te openen.")

        if not lessons:
            st.info("Je hebt nog geen lessen opgeslagen.")
        else:
            search_col, count_col = st.columns([4, 1.5])
            with search_col:
                search_query = st.text_input(
                    "Zoeken",
                    placeholder="🔍 Zoek op lesnaam, onderwerp, vak of klas...",
                    label_visibility="collapsed",
                    key="lessons_search_filter"
                )
            with count_col:
                st.caption(f"📁 Totaal: **{len(lessons)} les(sen)**")

            if search_query.strip():
                q = search_query.strip().lower()
                filtered_lessons = [
                    l for l in lessons
                    if q in (l.get("title") or "").lower()
                    or q in (l.get("topic") or "").lower()
                    or q in (l.get("subject") or "").lower()
                    or q in (l.get("class_name") or "").lower()
                ]
            else:
                filtered_lessons = lessons

            if not filtered_lessons:
                st.info("Geen lessen gevonden die overeenkomen met je zoekopdracht.")

            for lesson in filtered_lessons:
                analysis_obj = get_ai_analysis(lesson)
                if analysis_obj and analysis_obj.get("total_score") is not None:
                    try:
                        tot = int(analysis_obj["total_score"])
                        scale_5 = max(1, min(5, round((tot / 60) * 5)))
                        stars_str = "⭐️" * scale_5 + "☆" * (5 - scale_5)
                        score_badge = f"{stars_str} ({tot}/60)"
                        has_score = True
                    except (ValueError, TypeError):
                        score_badge = "⭐️ Geanalyseerd"
                        has_score = True
                else:
                    score_badge = "⚪ Geen AI-score"
                    has_score = False

                with st.container(border=True):
                    col_info, col_btn = st.columns([5.5, 1.2])
                    with col_info:
                        st.markdown(f"#### {lesson.get('title') or 'Naamloze les'}")
                        
                        badges = []
                        if lesson.get("subject"):
                            badges.append(f'<span class="detail-badge" style="font-size:0.78rem; padding:2px 8px;">📚 {html.escape(str(lesson["subject"]))}</span>')
                        if lesson.get("class_name"):
                            badges.append(f'<span class="detail-badge" style="font-size:0.78rem; padding:2px 8px;">👥 Klas: {html.escape(str(lesson["class_name"]))}</span>')
                        if lesson.get("date"):
                            badges.append(f'<span class="detail-badge" style="font-size:0.78rem; padding:2px 8px;">📅 {html.escape(str(lesson["date"]))}</span>')
                        if lesson.get("duration"):
                            badges.append(f'<span class="detail-badge" style="font-size:0.78rem; padding:2px 8px;">⏱️ {lesson["duration"]} min.</span>')
                        
                        if has_score:
                            badges.append(f'<span class="detail-badge" style="font-size:0.78rem; padding:2px 8px; background: rgba(251, 191, 36, 0.12); border-color: rgba(251, 191, 36, 0.35); color: #fbbf24;">{score_badge}</span>')
                        else:
                            badges.append(f'<span class="detail-badge" style="font-size:0.78rem; padding:2px 8px; color: #9ca3af; border-color: rgba(255,255,255,0.12);">{score_badge}</span>')

                        badges_html = f'<div style="display: flex; flex-wrap: wrap; gap: 6px; margin-top: 6px;">{" ".join(badges)}</div>'
                        st.markdown(badges_html, unsafe_allow_html=True)
                        
                    with col_btn:
                        st.write("")
                        st.write("")
                        if st.button("Bekijk les ➔", key=f"open_lesson_{lesson['id']}", use_container_width=True, type="secondary"):
                            st.session_state["selected_lesson_id"] = lesson["id"]
                            st.rerun()

    # 2. DETAILVENSTER VOOR DE GESELECTEERDE LES
    else:
        lesson = next((l for l in lessons if l["id"] == selected_id), None)
        if not lesson:
            st.session_state["selected_lesson_id"] = None
            st.warning("Deze les kon niet gevonden worden.")
            st.rerun()

        edit_key, delete_key = f"editing_lesson_{lesson['id']}", f"deleting_lesson_{lesson['id']}"
        st.session_state.setdefault(edit_key, False)
        st.session_state.setdefault(delete_key, False)

        clean_title = "".join(c for c in (lesson.get('title') or "les") if c.isalnum() or c in (' ', '_', '-')).strip().replace(' ', '_')

        top_l, top_r = st.columns([2.5, 4.5])
        with top_l:
            if st.button("⬅️ Terug naar lessen", key="back_to_overview_btn", type="secondary"):
                st.session_state["selected_lesson_id"] = None
                st.rerun()
        with top_r:
            if not st.session_state[edit_key]:
                b_cols = st.columns([1.3, 1.3, 1.4, 1.4])
                with b_cols[0]:
                    if st.button("✏️ Bewerk", key=f"edit_{lesson['id']}", help="Les bewerken", use_container_width=True):
                        st.session_state[edit_key] = True
                        st.rerun()
                with b_cols[1]:
                    if st.button("🗑️ Wis", key=f"delete_{lesson['id']}", help="Les definitief verwijderen", use_container_width=True):
                        st.session_state[delete_key] = True
                        st.rerun()
                with b_cols[2]:
                    _, cur_phases_db = get_lesson(lesson["id"])
                    cur_analysis = st.session_state.get(f"analysis_{lesson['id']}") or get_ai_analysis(lesson)
                    if DOCX_AVAILABLE:
                        docx_buffer = export_lesson_to_docx(lesson, cur_phases_db, cur_analysis)
                        if docx_buffer:
                            st.download_button(label="📄 Word", data=docx_buffer, file_name=f"Lesfiche_{clean_title}_{lesson.get('date') or 'export'}.docx", mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document", key=f"dl_docx_{lesson['id']}", help="Download als Word-document (.docx)", use_container_width=True)
                    else: st.caption("`python-docx`")
                with b_cols[3]:
                    txt_data = export_lesson_to_text(lesson, cur_phases_db, cur_analysis)
                    st.download_button(label="📥 TXT", data=txt_data, file_name=f"Lesvoorbereiding_{clean_title}_{lesson.get('date') or 'export'}.txt", mime="text/plain", key=f"dl_txt_{lesson['id']}", help="Download als tekstbestand (.txt)", use_container_width=True)

        if not st.session_state[edit_key]:
            versions = get_lesson_versions(lesson["id"])
            active_lesson = dict(lesson)
            active_phases_db = None
            active_analysis = None
            is_viewing_old_version = False
            current_version_str = "Actuele versie"

            if versions:
                with st.container(border=True):
                    v_col1, v_col2 = st.columns([3, 2])
                    with v_col1:
                        v_options = ["Huidige versie (Actueel)"] + [f"v{v['version_number']} — {v['created_at'][:16].replace('T', ' ')}" for v in versions]
                        chosen_v = st.selectbox("🕒 Versiegeschiedenis:", options=v_options, key=f"ver_select_{lesson['id']}")

                    if chosen_v != "Huidige versie (Actueel)":
                        is_viewing_old_version = True
                        v_idx = v_options.index(chosen_v) - 1
                        v_selected_record = versions[v_idx]
                        v_snapshot = json.loads(v_selected_record["snapshot_json"])

                        active_lesson = v_snapshot
                        active_lesson["id"] = lesson["id"]
                        active_phases_db = v_snapshot.get("phases", [])
                        current_version_str = f"Versie {v_selected_record['version_number']} (Historisch)"
                        if v_snapshot.get("analysis_json"):
                            try: active_analysis = json.loads(v_snapshot["analysis_json"]) if isinstance(v_snapshot["analysis_json"], str) else v_snapshot["analysis_json"]
                            except Exception: active_analysis = None

                        with v_col2:
                            st.write("")
                            st.write("")
                            if st.button("⏪ Herstel als actieve les", key=f"restore_btn_{v_selected_record['id']}", type="primary"):
                                restore_lesson_version(lesson["id"], v_selected_record["id"])
                                st.success(f"Versie {v_selected_record['version_number']} is teruggezet als actieve les!")
                                st.rerun()
                        st.info(f"Je bekijkt een historische momentopname: **v{v_selected_record['version_number']}** van **{v_selected_record['created_at'][:16].replace('T', ' ')}**.")
                    else:
                        current_version_str = f"Versie {versions[0]['version_number'] + 1} (Actueel)"

            if not is_viewing_old_version:
                _, active_phases_db = get_lesson(lesson["id"])
                active_analysis = st.session_state.get(f"analysis_{lesson['id']}") or get_ai_analysis(lesson)

            if st.session_state[delete_key]:
                with st.container(border=True):
                    st.warning(f"⚠️ Weet je zeker dat je de les '{lesson['title']}' definitief wilt verwijderen?")
                    confirm_col, cancel_col = st.columns(2)
                    with confirm_col:
                        if st.button("⚠️ Ja, definitief verwijderen", key=f"confirm_delete_{lesson['id']}", type="primary", use_container_width=True):
                            delete_lesson(lesson["id"])
                            st.session_state["selected_lesson_id"] = None
                            for k in list(st.session_state.keys()):
                                if str(lesson['id']) in k: st.session_state.pop(k, None)
                            st.rerun()
                    with cancel_col:
                        if st.button("Annuleren", key=f"cancel_delete_{lesson['id']}", use_container_width=True):
                            st.session_state[delete_key] = False
                            st.rerun()

            with st.container(border=True):
                st.markdown(f"## 📖 {active_lesson.get('title') or 'Naamloze les'}")
                if active_lesson.get('topic'):
                    st.markdown(f"**Onderwerp:** *{active_lesson['topic']}*")
                
                badges_html = f"""
                <div style="display: flex; flex-wrap: wrap; gap: 8px; margin-top: 12px;">
                    <span class="detail-badge">📚 {html.escape(str(active_lesson.get('subject') or 'Nederlands'))}</span>
                    <span class="detail-badge">🏛️ {html.escape(str(active_lesson.get('network') or 'GO!'))}</span>
                    <span class="detail-badge">👥 Klas: {html.escape(str(active_lesson.get('class_name') or '—'))}</span>
                    <span class="detail-badge">🎓 {html.escape(str(active_lesson.get('grade') or '3'))}e graad ({html.escape(str(active_lesson.get('study_direction') or 'Dubbele finaliteit'))})</span>
                    <span class="detail-badge">⏱️ {active_lesson.get('duration') or 50} min.</span>
                    <span class="detail-badge">📅 {html.escape(str(active_lesson.get('date') or 'Geen datum'))}</span>
                </div>
                """
                st.markdown(badges_html, unsafe_allow_html=True)

            with st.container(border=True):
                st.markdown("#### 🎯 Algemene lesdoelen")
                st.write(active_lesson.get("learning_objectives") or "Geen algemene lesdoelen opgegeven.")

            c_goals = normalize_multi_value(active_lesson.get("curriculum_goals"))
            if c_goals:
                with st.expander(f"🎯 Gekoppelde Leerplandoelen ({len(c_goals)})"):
                    for g_str in c_goals:
                        st.markdown(f"- `{g_str}`")

            if active_lesson.get("notes"):
                with st.container(border=True):
                    st.markdown("#### 📝 Algemene opmerkingen")
                    st.write(active_lesson.get("notes"))
                    
            st.divider()
            st.subheader("📍 Waar zijn we geëindigd?")
            st.caption("Houd bij waar je de les hebt afgerond om makkelijk verder te kunnen bij de volgende les.")

            with st.form(key=f"progress_form_{lesson['id']}", clear_on_submit=True):
                col_input, col_btn = st.columns([4, 1])
                with col_input:
                    new_eindpunt = st.text_input(
                        "Deze les geëindigd bij:",
                        placeholder="Bijvoorbeeld: Oefening 5, pagina 44",
                        label_visibility="collapsed"
                    )
                with col_btn:
                    submit_progress = st.form_submit_button("💾 Opslaan", use_container_width=True)
                
                if submit_progress and new_eindpunt.strip():
                    save_lesson_progress(lesson["id"], new_eindpunt.strip())
                    st.success(f"Opgeslagen: {new_eindpunt}")
                    st.rerun()
                elif submit_progress and not new_eindpunt.strip():
                    st.error("Voer eerst een eindpunt in.")

            progress_history = get_lesson_progress(lesson["id"])
            if progress_history:
                st.markdown(f"#### 📋 Geschiedenis ({len(progress_history)} items)")
                
                for item in progress_history:
                    try:
                        ts = datetime.fromisoformat(item["timestamp"])
                        formatted_date = ts.strftime("%d-%m-%Y om %H:%M")
                    except:
                        formatted_date = item["timestamp"]
                    
                    prog_col1, prog_col2 = st.columns([6, 1])
                    with prog_col1:
                        st.markdown(f"**{item['eindpunt']}**")
                        st.caption(f" Opgeslagen op {formatted_date}")
                    with prog_col2:
                        delete_key = f"delete_progress_{item['id']}"
                        if st.button("❌", key=delete_key, help="Verwijder dit eindpunt"):
                            st.session_state[f"confirm_delete_progress_{item['id']}"] = True
                            st.rerun()
                    
                    if st.session_state.get(f"confirm_delete_progress_{item['id']}", False):
                        with st.container(border=True):
                            st.warning(f"⚠️ Weet je zeker dat je '{item['eindpunt']}' wilt verwijderen?")
                            confirm_col1, confirm_col2 = st.columns(2)
                            with confirm_col1:
                                if st.button("✅ Ja, verwijder", key=f"confirm_yes_progress_{item['id']}", type="primary", use_container_width=True):
                                    delete_lesson_progress(item["id"])
                                    st.session_state[f"confirm_delete_progress_{item['id']}"] = False
                                    st.success("Verwijderd!")
                                    st.rerun()
                            with confirm_col2:
                                if st.button("❌ Annuleren", key=f"confirm_no_progress_{item['id']}", use_container_width=True):
                                    st.session_state[f"confirm_delete_progress_{item['id']}"] = False
                                    st.rerun()
                    st.divider()
            else:
                st.info("Nog geen eindpunten opgeslagen voor deze les.")

            st.markdown("---")
            st.markdown("##### 📊 PowerPoint Presentatie voor op het Digibord (Bouwsteen 2)")
            st.caption("Genereer een 16:9 presentatie. Slide 1 is de welkomstslide, Slide 2 toont de succescriteria en de lesfasen volgen. Leraarinstructies staan uitsluitend in de sprekersnotities.")
            
            ppt_col1, ppt_col2 = st.columns([2, 3])
            with ppt_col1:
                if st.button("📊 Genereer PowerPoint", key=f"gen_ppt_{lesson['id']}", help="Laat de AI succescriteria formuleren en genereer de presentatie"):
                    
                    ppt_progress = st.progress(10, text="Stap 1/2: AI vertaalt doelen naar succescriteria (Bouwsteen 2)...")
                    try:
                        s_goals = generate_student_goals(active_lesson.get("learning_objectives", ""), provider=st.session_state.ai_provider, model_name=st.session_state['gekozen_model'])
                        st.session_state[f"student_goals_{lesson['id']}"] = s_goals
                        
                        ppt_progress.progress(60, text="Stap 2/2: Presentatie dia's ontwerpen en vullen...")
                        ppt_buf = export_lesson_to_pptx(active_lesson, active_phases_db, student_goals=s_goals)
                        st.session_state[f"pptx_buf_{lesson['id']}"] = ppt_buf
                        
                        ppt_progress.progress(100, text="✅ PowerPoint-presentatie succesvol aangemaakt!")
                    except Exception as e:
                        ppt_progress.empty() # Verwijder progress bar bij fout
                        st.error(f"Fout bij genereren PowerPoint: {e}")
            with ppt_col2:
                if f"pptx_buf_{lesson['id']}" in st.session_state and st.session_state[f"pptx_buf_{lesson['id']}"]:
                    st.download_button(
                        label="📥 Download PowerPoint (.pptx)",
                        data=st.session_state[f"pptx_buf_{lesson['id']}"],
                        file_name=f"Lespresentatie_{clean_title}_{active_lesson.get('date') or 'export'}.pptx",
                        mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                        key=f"dl_pptx_{lesson['id']}",
                        type="primary"
                    )
                elif not PPTX_AVAILABLE:
                    st.caption("ℹ️ *Installeer `python-pptx` om PowerPoint-exports in te schakelen: `pip install python-pptx`*")

            st.markdown("---")
            with st.expander("📋 Smartschool Lesfiche (Kopieerbaar voor Schoolagenda)"):
                st.caption("Kopieer onderstaande tekst en plak deze direct in je Smartschool lesfiche of agenda:")
                smartschool_txt = format_for_smartschool(active_lesson, active_phases_db)
                st.code(smartschool_txt, language="markdown")

            display_lesson_structure(active_phases_db)
            st.divider()

            if not is_viewing_old_version:
                analyzing_key = f"is_analyzing_{lesson['id']}"
                if analyzing_key not in st.session_state:
                    st.session_state[analyzing_key] = False

                if not st.session_state[analyzing_key]:
                    if st.button("🔍 Analyseer deze les", key=f"analyze_lesson_{lesson['id']}", type="primary"):
                        st.session_state[analyzing_key] = True
                        st.rerun()
                else:
                    col_status, col_abort = st.columns([3, 1])
                    with col_abort:
                        if st.button("🛑 Analyse stoppen", key=f"abort_analysis_{lesson['id']}", type="secondary", use_container_width=True):
                            st.session_state[analyzing_key] = False
                            st.warning("⚠️ Analyse geannuleerd.")
                            st.rerun()
                    with col_status:
                        analysis_progress = st.progress(10, text=f"🤖 De AI analyseert je les met '{st.session_state['gekozen_model']}'... (Stap 1/3: Voorbereiden)")

                    try:
                        lesson_text = build_lesson_text(active_lesson, active_phases_db)
                        
                        analysis_progress.progress(30, text="Stap 2/3: AI voert de didactische analyse uit over de 12 bouwstenen (dit kan even duren)...")
                        analysis = analyze_lesson(lesson_text, provider=st.session_state.ai_provider, model_name=st.session_state['gekozen_model'])
                        analysis["analyzed_at"] = datetime.now().strftime("%d-%m-%Y om %H:%M")
                        
                        analysis_progress.progress(70, text="Stap 3/3: AI genereert Spaced Retrieval oefeningen...")
                        try:
                            spaced_exs = generate_spaced_retrieval(lesson_text, provider=st.session_state.ai_provider, model_name=st.session_state['gekozen_model'])
                            analysis["spaced_retrieval"] = spaced_exs
                        except Exception:
                            analysis["spaced_retrieval"] = []

                        analysis_progress.progress(95, text="Resultaten opslaan...")
                        save_ai_analysis(lesson["id"], analysis)
                        st.session_state[f"analysis_{lesson['id']}"] = analysis
                        st.session_state[analyzing_key] = False
                        
                        analysis_progress.progress(100, text="✅ Analyse succesvol afgerond!")
                        st.rerun()
                    except Exception as error:
                        st.session_state[analyzing_key] = False
                        st.error("Er ging iets mis bij de AI-analyse.")
                        st.exception(error)

            if active_analysis: 
                show_ai_analysis(active_analysis, lesson_obj=active_lesson, version_str=current_version_str)

        else:
            st.subheader("✏️ Les bewerken")
            current_lesson, current_phases = get_lesson(lesson["id"])
            editor_key = f"edit_phases_{lesson['id']}"

            if editor_key not in st.session_state: st.session_state[editor_key] = db_phases_to_editor(current_phases)
            existing = st.session_state[editor_key]

            st.markdown("### Basisinformatie")
            col1, col2 = st.columns(2)
            with col1:
                edit_title = st.text_input("Lesnaam *", value=current_lesson.get("title", ""), key=f"edit_title_{lesson['id']}")
                edit_class_name = st.text_input("Klas", value=current_lesson.get("class_name") or "", key=f"edit_class_{lesson['id']}")
                edit_grade = st.text_input("Leerjaar/Graad", value=current_lesson.get("grade") or "3", key=f"edit_grade_{lesson['id']}")
            with col2:
                edit_study_direction = st.text_input("Studierichting/Finaliteit", value=current_lesson.get("study_direction") or "Dubbele finaliteit", key=f"edit_direction_{lesson['id']}")
                try: current_date = datetime.strptime(current_lesson.get("date", ""), "%Y-%m-%d").date()
                except (ValueError, TypeError): current_date = date.today()
                edit_lesson_date = st.date_input("Datum", value=current_date, key=f"edit_date_{lesson['id']}")
                edit_duration = st.number_input("Totale lesduur (minuten)", min_value=1, max_value=240, value=int(current_lesson.get("duration") or 50), step=5, key=f"edit_duration_{lesson['id']}")

            edit_topic = st.text_input("Onderwerp", value=current_lesson.get("topic") or "", key=f"edit_topic_{lesson['id']}")
            edit_learning_objectives = st.text_area("Algemene leerdoelen", value=current_lesson.get("learning_objectives") or "", key=f"edit_objectives_{lesson['id']}")

            available_goals_edit = load_curriculum_goals(net=current_lesson.get("network") or "GO!", graad=current_lesson.get("grade") or "3", finaliteit=current_lesson.get("study_direction") or "Dubbele finaliteit", vak=current_lesson.get("subject") or "Nederlands")
            options_edit = [f"{g['code']} - {g['omschrijving']} ({g.get('cluster', '')})" for g in available_goals_edit]
            saved_goals = normalize_multi_value(current_lesson.get("curriculum_goals"))
            edit_selected_goals = st.multiselect("🎯 Gekoppelde Leerplandoelen", options=options_edit, default=[item for item in saved_goals if item in options_edit], key=f"edit_goals_{lesson['id']}")

            edit_notes = st.text_area("Algemene opmerkingen", value=current_lesson.get("notes") or "", key=f"edit_notes_{lesson['id']}")
            st.divider()
            st.markdown("### Lesvoorbereiding")
            edit_phases = render_edit_structure(existing, lesson["id"])
            st.divider()
            st.markdown("### ⏱️ Tijdcontrole")
            show_duration_check(edit_duration, edit_phases)

            save_changes = st.button("💾 Wijzigingen opslaan", key=f"save_edit_{lesson['id']}", type="primary", use_container_width=True)
            cancel_edit = st.button("Annuleren", key=f"cancel_edit_{lesson['id']}")

            if cancel_edit:
                keys_to_clear = [k for k in st.session_state.keys() if f"_{lesson['id']}" in k and (k.startswith("edit_") or k.startswith("editing_"))]
                for k in keys_to_clear: st.session_state.pop(k, None)
                st.session_state[edit_key] = False
                st.rerun()

            if save_changes:
                if not edit_title.strip(): st.error("Geef eerst een naam voor de les.")
                elif not edit_learning_objectives.strip(): st.error("Voeg minstens één algemeen leerdoel toe.")
                else:
                    update_lesson(lesson["id"], edit_title, edit_class_name, edit_grade, edit_study_direction, edit_lesson_date.isoformat(), edit_duration, edit_topic, edit_learning_objectives, edit_notes, edit_phases, network=current_lesson.get("network") or "GO!", subject=current_lesson.get("subject") or "Nederlands", curriculum_goals=edit_selected_goals)
                    keys_to_clear = [k for k in st.session_state.keys() if f"_{lesson['id']}" in k and (k.startswith("edit_") or k.startswith("editing_"))]
                    for k in keys_to_clear: st.session_state.pop(k, None)
                    st.session_state[edit_key] = False
                    st.success("De les is bijgewerkt en de vorige versie is bewaard!")
                    st.rerun()

# ============================================================
# PAGINA: 🎯 DOELEN TRACKER (CURRICULUM COVERAGE)
# ============================================================
elif page == "🎯 Doelen Tracker":
    st.title("🎯 Leerplandoelen Tracker")
    st.caption("Krijg in één oogopslag inzicht in de dekking van het officiële leerplan over al jouw opgeslagen lessen heen.")
    st.markdown("#### 🎯 Selecteer Leerplan")
    t_col1, t_col2, t_col3, t_col4 = st.columns(4)
    with t_col1: tr_network = st.selectbox("Onderwijsnet", ["GO!", "Katholiek Onderwijs", "OVSG"], index=0, key="tr_net")
    with t_col2: tr_grade = st.selectbox("Graad / Leerjaar", ["1", "2", "3", "5", "6"], index=2, key="tr_grd")
    with t_col3: tr_finality = st.selectbox("Finaliteit", ["Dubbele finaliteit", "Doorstroom", "Arbeidsmarkt"], index=0, key="tr_fin")
    with t_col4: tr_subject = st.selectbox("Vak", ["Nederlands", "Wiskunde", "Geschiedenis", "Engels"], index=0, key="tr_subj")

    available_goals = load_curriculum_goals(net=tr_network, graad=tr_grade, finaliteit=tr_finality, vak=tr_subject)

    if not available_goals: st.warning(f"Geen leerplanbestand gevonden in `data/leerplannen/` voor **{tr_network} - {tr_grade}e graad {tr_finality} ({tr_subject})**.")
    else:
        all_lessons = get_lessons()
        matching_lessons = [l for l in all_lessons if (not l.get("network") or l.get("network") == tr_network) and (not l.get("subject") or l.get("subject") == tr_subject) and (not l.get("grade") or str(l.get("grade")) == str(tr_grade))]
        covered_map = {g["code"]: [] for g in available_goals}
        for l in matching_lessons:
            raw_c_goals = normalize_multi_value(l.get("curriculum_goals"))
            for g in available_goals:
                code = g["code"]
                if any(code in g_str for g_str in raw_c_goals): covered_map[code].append(l)

        total_goals = len(available_goals)
        covered_count = sum(1 for code, les_list in covered_map.items() if len(les_list) > 0)
        uncovered_count = total_goals - covered_count
        coverage_pct = int((covered_count / total_goals) * 100) if total_goals > 0 else 0

        st.divider()
        st.markdown("#### 📈 Voortgang & Dekkingsgraad")
        m_col1, m_col2, m_col3, m_col4 = st.columns(4)
        with m_col1: tr_coverage = st.metric("Dekkingsgraad", f"{coverage_pct}%")
        with m_col2: tr_coverage_count = st.metric("Behandelde Doelen", f"{covered_count} / {total_goals}")
        with m_col3: tr_uncovered = st.metric("Nog te Plannen", f"{uncovered_count}")
        with m_col4: tr_lessons = st.metric("Geanalyseerde Lessen", len(matching_lessons))
        st.progress(coverage_pct / 100.0)

        st.divider()
        f_col1, f_col2 = st.columns([2, 2])
        with f_col1: view_status = st.radio("Filter op status:", ["Alle doelen", "Alleen behandeld (Groen)", "Alleen nog te plannen (Grijs)"], horizontal=True)
        with f_col2: view_type = st.radio("Filter op type competentie:", ["Alle competenties", "Vakgerelateerd", "Vakoverschrijdend"], horizontal=True)

        filtered_goals = []
        for g in available_goals:
            is_covered = len(covered_map.get(g["code"], [])) > 0
            if view_status == "Alleen behandeld (Groen)" and not is_covered: continue
            if view_status == "Alleen nog te plannen (Grijs)" and is_covered: continue
            g_type = g.get("type", "").lower()
            if view_type == "Vakgerelateerd" and "vakgerelateerd" not in g_type: continue
            if view_type == "Vakoverschrijdend" and "vakoverschrijdend" not in g_type: continue
            filtered_goals.append((g, is_covered, covered_map.get(g["code"], [])))

        st.markdown(f"**Weergave:** {len(filtered_goals)} doelen")
        for g, is_covered, lesson_hits in filtered_goals:
            card_class = "tracker-card-covered" if is_covered else "tracker-card-uncovered"
            status_icon = "🟢 **BEHANDELD**" if is_covered else "⚪ **NOG NIET BEHANDELD**"
            g_type_label = "Vakcompetentie" if g.get("type") == "vakgerelateerd" else "Vakoverschrijdend"
            cluster_info = f" · *{g.get('cluster', '')}*" if g.get('cluster') else ""

            st.markdown(f'<div class="tracker-card {card_class}"><div style="display: flex; justify-content: space-between; align-items: center; margin-bottom: 6px;"><span style="font-weight: 700; font-size: 1.05rem; color: #60a5fa;">{g["code"]}</span><span style="font-size: 0.85rem;">{status_icon} ({len(lesson_hits)}x)</span></div><div style="font-size: 0.95rem; line-height: 1.45; margin-bottom: 6px;">{html.escape(g["omschrijving"])}</div><div style="font-size: 0.8rem; color: #a1a1aa;">🏷️ <strong>{g_type_label}</strong>{cluster_info}</div></div>', unsafe_allow_html=True)
            if is_covered:
                with st.expander(f"📚 Bekijk {len(lesson_hits)} gekoppelde les(sen) voor {g['code']}"):
                    for l_hit in lesson_hits: st.markdown(f"- **{l_hit['title']}** ({l_hit['date'] or 'Geen datum'}) — *Klas: {l_hit.get('class_name') or '—'}*")

elif page == "🧱 De 12 bouwstenen":
    st.title("📖 De 12 bouwstenen")
    st.write("De didactische principes waarop de coach zijn analyses baseert.")
    for principle in PRINCIPLES:
        with st.expander(f"{principle['emoji']} {principle['number']}. {principle['name']}"):
            st.write(principle["description"])

st.sidebar.caption("Wijze Lessen Coach · lokale & online versie")